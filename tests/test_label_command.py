"""The `label` command end to end (commands/label.py).

Runs label on small sparse alignments and reads the XLSX workbook back:
argument errors, sheet layout and percentages, subset-specific residues and
the gmax outside background, output filenames, and the job snapshot. The
analysis helpers are covered in test_label_analysis.
"""

import io
import os
import sys
import tempfile
import threading
import time
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

import numpy as np
import openpyxl


PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_DIR = os.path.join(PROJECT_ROOT, "src")
if SRC_DIR not in sys.path:
    sys.path.insert(0, SRC_DIR)

import Alignment_Manager
from commands import label
from desktop import Desktop_App as application_windows
from tests.sparse_alignment import sparse_alignment
from utilities.Job_Cancellation import cancel_requested


class AlignmentStub:
    """Alignment_Manager state for label: a sparse alignment plus a fixed mapping."""

    def __init__(self, sequences=None):
        sequences = ["A", "A", "C"] if sequences is None else list(sequences)
        self.aln = sparse_alignment(
            (f"node{index}", sequence) for index, sequence in enumerate(sequences)
        )
        self.col_to_label = {0: "1"}
        self.label_to_col = {"1": 0}
        self.has_reference = True
        self.resolved_ref_full = "node0"
        self.viewer_to_aln = np.arange(len(sequences), dtype=int)


class ImmediateScheduler:
    def __init__(self, viewer):
        self.viewer = viewer

    def is_output_path_reserved(self, _path):
        return False

    def enqueue(self, **job):
        result = job["worker"](job["payload"])
        self.viewer.console_text.text = str(result["message"])
        reveal_directory = result.get("reveal_directory")
        if reveal_directory:
            application_windows.open_in_file_manager(reveal_directory)
        return 1


class CapturingScheduler:
    def __init__(self):
        self.job = None

    def is_output_path_reserved(self, _path):
        return False

    def enqueue(self, **job):
        self.job = job
        return 1


def find_row(worksheet, first_cell_value):
    for row_idx in range(1, worksheet.max_row + 1):
        if worksheet.cell(row=row_idx, column=1).value == first_cell_value:
            return row_idx
    raise AssertionError(f"Could not find row {first_cell_value!r}")


class LabelWorkbookPercentTests(unittest.TestCase):
    @staticmethod
    def make_viewer(sequences, cluster_labels, group_labels):
        headers = [f"node{index}" for index in range(len(sequences))]
        return SimpleNamespace(
            alignment=AlignmentStub(sequences),
            active_reference="node0",
            full_headers=headers,
            n_nodes=len(headers),
            cluster_labels=(
                None
                if cluster_labels is None
                else np.asarray(cluster_labels, dtype=int)
            ),
            group_labels=[set(groups) for groups in group_labels],
            console_text=SimpleNamespace(text=""),
            last_cluster_params=None,
        )

    def run_label(
        self, directory, args, viewer=None, identity_weights=None, real_colors=False,
        scheduler=None, after_run=None,
    ):
        """Run label on a small alignment; its job runs at once, or on ``scheduler`` if one is given.

        ``after_run(viewer)`` is called while the patches still hold, so that a job a
        scheduler's thread runs meanwhile sees them.
        """
        if viewer is None:
            viewer = SimpleNamespace(
                alignment=AlignmentStub(),
                active_reference="node0",
                full_headers=["node0", "node1", "node2", "unaligned"],
                n_nodes=4,
                cluster_labels=np.array([0, 0, 1, -1]),
                group_labels=[{"GroupA"}, set(), set(), set()],
                console_text=SimpleNamespace(text=""),
                last_cluster_params=None,
            )
        viewer.background_job_scheduler = scheduler if scheduler is not None else ImmediateScheduler(viewer)
        viewer_to_aln = getattr(viewer.alignment, "viewer_to_aln", None)
        if viewer_to_aln is None or np.asarray(viewer_to_aln).shape != (
            len(viewer.full_headers),
        ):
            viewer_to_aln = np.full(len(viewer.full_headers), -1, dtype=int)
            aligned_count = min(len(viewer.alignment.aln), len(viewer.full_headers))
            viewer_to_aln[:aligned_count] = np.arange(aligned_count)
        viewer_to_aln = np.asarray(viewer_to_aln, dtype=int)
        if identity_weights is None:
            identity_weights = np.array([0.5, 0.5, 1.0])
        if real_colors:
            color_map_patch = mock.patch.object(
                label.cluster_cmd,
                "get_cluster_color_map",
                wraps=label.cluster_cmd.get_cluster_color_map,
            )
        else:
            color_map_patch = mock.patch.object(
                label.cluster_cmd,
                "get_cluster_color_map",
                return_value={0: (1.0, 0.0, 0.0), 1: (0.0, 1.0, 0.0)},
            )

        with mock.patch.object(label, "CLUSTER_LABEL_DIRECTORY", directory), \
                mock.patch.object(label.cfg, "NODE_FASTA_FILE", "nodes.fasta"), \
                mock.patch.object(label.cfg, "INPUT_HDF5", "network.h5"), \
                mock.patch.object(label.cfg, "MSA_FILE", "alignment.fasta"), \
                mock.patch.object(
                    label.Command_Engine,
                    "get_alignment_mapping",
                    return_value=(
                        viewer_to_aln,
                        np.flatnonzero(viewer_to_aln >= 0),
                    ),
                ), \
                color_map_patch, \
                mock.patch.object(
                    label.Alignment_Profiles,
                    "calculate_identity_weights",
                    return_value=np.asarray(identity_weights, dtype=float),
                ) as identity_weight_mock, \
                mock.patch.object(application_windows, "open_in_file_manager"):
            label.run(viewer, args)
            if after_run is not None:
                after_run(viewer)
        viewer.identity_weight_mock = identity_weight_mock
        return viewer

    def test_percent_column_uses_total_network_nodes_in_both_sheets(self):
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(directory, [])

            output_path = next(Path(directory).glob("Label_Output_*.xlsx"))
            workbook = openpyxl.load_workbook(output_path)

            for sheet_name in ("Subset Stats", "Occupancy Stats"):
                worksheet = workbook[sheet_name]
                header_row = find_row(worksheet, "Subset Name")
                cluster_row = find_row(worksheet, "Cluster 0")
                group_row = find_row(worksheet, "Group GroupA")
                global_row = find_row(worksheet, "Global Stats")

                self.assertEqual(worksheet.freeze_panes, "C1")
                self.assertEqual(worksheet.cell(header_row, 2).value, "Percent")
                self.assertEqual(worksheet.cell(header_row, 3).value, "Count")
                self.assertEqual(worksheet.cell(header_row, 4).value, "Hex Color")
                self.assertEqual(worksheet.cell(header_row, 10).value, "#1")
                self.assertEqual(worksheet.cell(cluster_row, 2).value, 0.5)
                self.assertEqual(worksheet.cell(group_row, 2).value, 0.25)
                self.assertEqual(worksheet.cell(global_row, 2).value, 0.75)
                self.assertEqual(worksheet.cell(cluster_row, 2).number_format, "0.00%")
                self.assertEqual(worksheet.column_dimensions["B"].width, 10.0)
                self.assertEqual(worksheet.cell(cluster_row, 4).fill.fill_type, "solid")
                self.assertEqual(worksheet.cell(cluster_row, 10).fill.fill_type, "solid")
            viewer = self.run_label(directory, ["second_report"])
            viewer.identity_weight_mock.assert_not_called()

    def test_custom_filename_is_used_and_xlsx_is_appended(self):
        with tempfile.TemporaryDirectory() as directory:
            viewer = self.run_label(
                directory,
                ["0.4", "0.9", "custom_label_report"],
            )

            output_path = Path(directory, "custom_label_report.xlsx")
            self.assertTrue(output_path.is_file())
            self.assertIn(str(output_path), viewer.console_text.text)
            self.assertEqual(
                [path for path in Path(directory).iterdir() if ".partial" in path.name],
                [],
            )

            workbook = openpyxl.load_workbook(output_path)
            metadata_sheet = workbook["Meta Data"]
            alignment_row = find_row(metadata_sheet, "Alignment Name")
            label_row = find_row(metadata_sheet, "Label Parameters")
            statistics_row = find_row(metadata_sheet, "Statistics")
            self.assertEqual(label_row, alignment_row + 1)
            self.assertEqual(statistics_row, label_row + 1)
            self.assertEqual(
                metadata_sheet.cell(label_row, 2).value,
                "gmax_outside=40%, cmin=90%, target=all",
            )
            self.assertEqual(
                metadata_sheet.cell(statistics_row, 2).value,
                "Aligned 3 of 4 | Excluded 1 | Effective 3",
            )
            for sheet_name in ("Subset Stats", "Occupancy Stats"):
                worksheet = workbook[sheet_name]
                self.assertIsNotNone(find_row(
                    worksheet,
                    "Statistics: Aligned 3 of 4 | Excluded 1 | Effective 3",
                ))
                self.assertIsNotNone(find_row(worksheet, "Global Conserved (>97%)"))

    def test_custom_filename_overwrites_existing_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory, "overwrite_report.xlsx")
            output_path.write_bytes(b"previous content")

            viewer = self.run_label(directory, ["overwrite_report"])

            self.assertIn(str(output_path), viewer.console_text.text)
            self.assertNotEqual(output_path.read_bytes(), b"previous content")
            workbook = openpyxl.load_workbook(output_path)
            self.assertIn("Meta Data", workbook.sheetnames)
            self.assertEqual(
                [path for path in Path(directory).iterdir() if ".partial" in path.name],
                [],
            )

    def test_automatic_output_created_during_render_is_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory, "automatic.xlsx")

            def save_and_create_competing_output(
                _workbook,
                target,
            ):
                # The workbook is saved into the partial file through label's cancellable file.
                target.write(b"rendered content")
                output_path.write_bytes(b"competing content")

            with mock.patch.object(
                label.logo_cmd,
                "_available_automatic_filename",
                return_value=(output_path.name, str(output_path)),
            ), mock.patch(
                "openpyxl.workbook.workbook.Workbook.save",
                autospec=True,
                side_effect=save_and_create_competing_output,
            ):
                viewer = self.run_label(directory, [])

            self.assertIn("already exists", viewer.console_text.text)
            self.assertEqual(output_path.read_bytes(), b"competing content")
            self.assertEqual(
                [path for path in Path(directory).iterdir() if ".partial" in path.name],
                [],
            )

    def test_bare_filename_is_allowed_after_keyword_parameters(self):
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(directory, ["cmin", "90%", "keyword_report"])

            self.assertTrue(Path(directory, "keyword_report.xlsx").is_file())

    def test_identity_keyword_forms_are_equivalent(self):
        for index, identity_value in enumerate(("0.9", "90", "90%")):
            with self.subTest(identity_value=identity_value), \
                    tempfile.TemporaryDirectory() as directory:
                viewer = self.run_label(
                    directory,
                    ["id", identity_value, f"identity_{index}"],
                )

                self.assertTrue(Path(directory, f"identity_{index}.xlsx").is_file())
                call_args = viewer.identity_weight_mock.call_args
                self.assertEqual(call_args.args[1], 0.9)
                self.assertTrue(call_args.kwargs["report_backend"])

    def test_third_positional_number_enables_identity_weighting(self):
        with tempfile.TemporaryDirectory() as directory:
            viewer = self.run_label(
                directory,
                ["0.4", "0.9", "100%", "weighted_report"],
            )

            output_path = Path(directory, "weighted_report.xlsx")
            workbook = openpyxl.load_workbook(output_path)
            metadata_rows = [
                (row[0].value, row[1].value)
                for row in workbook["Meta Data"].iter_rows(min_col=1, max_col=2)
            ]
            metadata = dict(metadata_rows)

            self.assertEqual(
                metadata["Statistics"],
                "Aligned 3 of 4 | Excluded 1 | Effective 2",
            )
            self.assertNotIn("Network Nodes", metadata)
            self.assertNotIn("Aligned Nodes", metadata)
            self.assertNotIn("Global Effective N", metadata)
            self.assertNotIn("Excluded Unaligned Nodes", metadata)
            self.assertNotIn("Identity Threshold", metadata)
            self.assertNotIn("Identity Backend", metadata)
            self.assertNotIn("Identity Threads", metadata)
            self.assertNotIn("Identity Fallback Reason", metadata)
            metadata_names = [name for name, _value in metadata_rows]
            alignment_index = metadata_names.index("Alignment Name")
            self.assertEqual(metadata_names[alignment_index + 1], "Label Parameters")
            self.assertEqual(metadata_names[alignment_index + 2], "Statistics")
            self.assertIn("identity=100%", metadata["Label Parameters"])
            self.assertNotIn("outside_pool", metadata["Label Parameters"])
            self.assertNotIn(
                "global_conservation_threshold",
                metadata["Label Parameters"],
            )
            viewer.identity_weight_mock.assert_called_once()

            for sheet_name in ("Subset Stats", "Occupancy Stats"):
                worksheet = workbook[sheet_name]
                self.assertIsNotNone(find_row(
                    worksheet,
                    "Statistics: Aligned 3 of 4 | Excluded 1 | Effective 2",
                ))
                header_row = find_row(worksheet, "Subset Name")
                global_row = find_row(worksheet, "Global Stats")
                cluster_row = find_row(worksheet, "Cluster 0")
                group_row = find_row(worksheet, "Group GroupA")

                self.assertEqual(worksheet.freeze_panes, "C1")
                self.assertEqual(
                    [worksheet.cell(header_row, column).value for column in range(1, 6)],
                    [
                        "Subset Name", "Effective Percent", "Percent",
                        "Effective N", "Count N",
                    ],
                )
                self.assertEqual(worksheet.cell(header_row, 6).value, "Hex Color")
                self.assertEqual(worksheet.cell(header_row, 12).value, "#1")
                self.assertEqual(worksheet.cell(global_row, 2).value, 1.0)
                self.assertEqual(worksheet.cell(global_row, 3).value, 0.75)
                self.assertEqual(worksheet.cell(global_row, 4).value, 2.0)
                self.assertEqual(worksheet.cell(global_row, 5).value, 3)
                self.assertEqual(worksheet.cell(cluster_row, 2).value, 0.5)
                self.assertEqual(worksheet.cell(cluster_row, 3).value, 0.5)
                self.assertEqual(worksheet.cell(cluster_row, 4).value, 1.0)
                self.assertEqual(worksheet.cell(cluster_row, 5).value, 2)
                self.assertEqual(worksheet.cell(group_row, 2).value, 0.25)
                self.assertEqual(worksheet.cell(group_row, 3).value, 0.25)
                self.assertEqual(
                    worksheet.cell(cluster_row, 2).number_format,
                    "0.00%",
                )
                self.assertEqual(
                    worksheet.cell(cluster_row, 3).number_format,
                    "0.00%",
                )
                self.assertEqual(
                    worksheet.cell(cluster_row, 4).number_format,
                    "0.00",
                )
                self.assertEqual(
                    worksheet.cell(cluster_row, 6).fill.fill_type,
                    "solid",
                )

    def test_identity_assignment_errors_do_not_write_workbooks(self):
        cases = (
            (["id"], "Missing numerical value for 'id'"),
            (["id", "0"], "outside the supported range"),
            (["id", "90%", "id", "80%"], "Duplicate assignment for 'id'"),
            (
                ["0.4", "0.9", "90%", "id", "80%"],
                "defined both positionally and via keyword",
            ),
            (["0.4", "0.9", "90%", "80%"], "Too many positional"),
        )
        for args, expected_error in cases:
            with self.subTest(args=args), tempfile.TemporaryDirectory() as directory:
                viewer = self.run_label(directory, args)

                self.assertIn(expected_error, viewer.console_text.text)
                self.assertEqual(list(Path(directory).glob("*.xlsx")), [])
                viewer.identity_weight_mock.assert_not_called()

    def test_numeric_filename_requires_and_accepts_xlsx_extension(self):
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(directory, ["0.4.xlsx"])

            self.assertTrue(Path(directory, "0.4.xlsx").is_file())

    def test_filename_must_be_final_and_old_keyword_form_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            viewer = self.run_label(directory, ["report", "cmin", "90%"])
            self.assertEqual(
                viewer.console_text.text,
                "Error: A custom output filename must be the final argument.",
            )
            self.assertEqual(list(Path(directory).glob("*.xlsx")), [])

        with tempfile.TemporaryDirectory() as directory:
            viewer = self.run_label(directory, ["filename", "old_style"])
            self.assertEqual(
                viewer.console_text.text,
                "Error: A custom output filename must be the final argument.",
            )
            self.assertEqual(list(Path(directory).glob("*.xlsx")), [])

    def test_gmin_is_not_user_configurable(self):
        with tempfile.TemporaryDirectory() as directory:
            viewer = self.run_label(directory, ["gmin", "90%"])

            self.assertEqual(
                viewer.console_text.text,
                "Error: gmin is fixed at 97% and cannot be set by the label command.",
            )
            self.assertEqual(list(Path(directory).glob("*.xlsx")), [])

    def test_each_failure_is_reported_once(self):
        with tempfile.TemporaryDirectory() as directory, mock.patch.object(
            label.Command_Engine, "command_failed"
        ) as command_failed:
            viewer = self.run_label(directory, ["gmin", "90%"])

        command_failed.assert_called_once()
        self.assertIs(command_failed.call_args.args[0], viewer)
        self.assertEqual(
            str(command_failed.call_args.args[1]),
            "Error: gmin is fixed at 97% and cannot be set by the label command.",
        )

    def test_missing_reference_fails_before_any_job_is_scheduled(self):
        viewer = self.make_viewer(["A", "A", "C"], [0, 0, 1], [set(), set(), set()])
        viewer.alignment.has_reference = False
        viewer.background_job_scheduler = CapturingScheduler()
        message = (
            "Error: No active alignment reference. Use 'reference <ID>' with a "
            "node present in the current MSA."
        )
        with mock.patch.object(label.Command_Engine, "command_failed") as command_failed, \
                redirect_stdout(io.StringIO()) as output:
            label.run(viewer, [])

        command_failed.assert_called_once()
        self.assertIs(command_failed.call_args.args[0], viewer)
        self.assertEqual(str(command_failed.call_args.args[1]), message)
        self.assertIsNone(viewer.background_job_scheduler.job)
        self.assertEqual(viewer.console_text.text, message)
        self.assertEqual(output.getvalue(), message + "\n")

    def test_global_stats_profile_divides_by_every_aligned_row(self):
        cases = {
            ("Y", "Y", "Y", "Y", "A", "A"): "Y  66.7% | A  33.3%",
            # The gap row counts in the denominator but adds no residue.
            ("Y", "Y", "Y", "Y", "A", "-"): "Y  66.7% | A  16.7%",
        }
        for index, (sequences, profile) in enumerate(cases.items()):
            viewer = self.make_viewer(
                list(sequences), [0, 0, 1, 1, 2, 2], [set() for _ in range(6)]
            )
            with self.subTest(sequences=sequences), \
                    tempfile.TemporaryDirectory() as directory:
                self.run_label(
                    directory,
                    ["gmax", "50%", "cmin", "100%", f"profile_{index}"],
                    viewer=viewer,
                )
                worksheet = openpyxl.load_workbook(
                    Path(directory, f"profile_{index}.xlsx")
                )["Subset Stats"]
                self.assertEqual(
                    worksheet.cell(find_row(worksheet, "Subset Name"), 10).value, "#1"
                )
                self.assertEqual(
                    worksheet.cell(find_row(worksheet, "Global Stats"), 10).value,
                    profile,
                )

    def test_global_conserved_row_needs_more_than_97_percent(self):
        # 33 of 34 rows is 97.06% Y; 32 of 33 is 96.97%.
        for y_rows, expected in ((33, ["Y1"]), (32, ["None"])):
            sequences = ["Y"] * y_rows + ["A"]
            viewer = self.make_viewer(
                sequences, [0] * len(sequences), [set() for _ in sequences]
            )
            with self.subTest(y_rows=y_rows), tempfile.TemporaryDirectory() as directory:
                self.run_label(directory, [f"conserved_{y_rows}"], viewer=viewer)
                workbook = openpyxl.load_workbook(
                    Path(directory, f"conserved_{y_rows}.xlsx")
                )
                for sheet_name in ("Subset Stats", "Occupancy Stats"):
                    worksheet = workbook[sheet_name]
                    row = find_row(worksheet, "Global Conserved (>97%)")
                    values = [
                        cell.value for cell in worksheet[row + 1] if cell.value is not None
                    ]
                    self.assertEqual(values, expected, sheet_name)

    def test_shared_union_reports_two_clusters_that_fail_individually(self):
        viewer = self.make_viewer(
            ["Y", "Y", "Y", "Y", "A", "A"],
            [0, 0, 1, 1, 2, 2],
            [set() for _ in range(6)],
        )
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(
                directory,
                ["gmax", "50%", "cmin", "100%", "shared_clusters"],
                viewer=viewer,
            )
            workbook = openpyxl.load_workbook(
                Path(directory, "shared_clusters.xlsx")
            )
            worksheet = workbook["Subset Stats"]
            header_row = find_row(worksheet, "Subset Name")
            self.assertEqual(worksheet.cell(header_row, 10).value, "#1")
            self.assertEqual(
                worksheet.cell(find_row(worksheet, "Cluster 0"), 10).value,
                "Y1",
            )
            self.assertEqual(
                worksheet.cell(find_row(worksheet, "Cluster 1"), 10).value,
                "Y1",
            )

    def test_default_mode_pools_clusters_and_groups_together(self):
        viewer = self.make_viewer(
            ["Y", "Y", "Y", "Y", "A"],
            [0, 0, 1, 1, 1],
            [set(), set(), {"SharedY"}, {"SharedY"}, set()],
        )
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(
                directory,
                ["gmax", "50%", "cmin", "100%", "cross_type"],
                viewer=viewer,
            )
            worksheet = openpyxl.load_workbook(
                Path(directory, "cross_type.xlsx")
            )["Subset Stats"]
            self.assertEqual(
                worksheet.cell(find_row(worksheet, "Cluster 0"), 10).value,
                "Y1",
            )
            self.assertEqual(
                worksheet.cell(find_row(worksheet, "Group SharedY"), 10).value,
                "Y1",
            )

    def test_cluster_aliases_exclude_group_results(self):
        for target in ("cluster", "clusters"):
            with self.subTest(target=target), tempfile.TemporaryDirectory() as directory:
                self.run_label(directory, [target, f"{target}_only"])
                worksheet = openpyxl.load_workbook(
                    Path(directory, f"{target}_only.xlsx")
                )["Subset Stats"]
                subset_names = {
                    worksheet.cell(row=row, column=1).value
                    for row in range(1, worksheet.max_row + 1)
                }

                self.assertIn("Cluster 0", subset_names)
                self.assertIn("Cluster 1", subset_names)
                self.assertNotIn("Group GroupA", subset_names)

    def test_group_aliases_exclude_cluster_results(self):
        for target in ("group", "groups"):
            with self.subTest(target=target), tempfile.TemporaryDirectory() as directory:
                self.run_label(directory, [target, f"{target}_only"])
                worksheet = openpyxl.load_workbook(
                    Path(directory, f"{target}_only.xlsx")
                )["Subset Stats"]
                subset_names = {
                    worksheet.cell(row=row, column=1).value
                    for row in range(1, worksheet.max_row + 1)
                }

                self.assertIn("Group GroupA", subset_names)
                self.assertNotIn("Cluster 0", subset_names)
                self.assertNotIn("Cluster 1", subset_names)

    def test_default_mode_uses_groups_when_clusters_are_unavailable(self):
        viewer = self.make_viewer(
            ["Y", "Y", "A"],
            None,
            [{"OnlyGroup"}, {"OnlyGroup"}, set()],
        )
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(directory, ["groups_available"], viewer=viewer)
            worksheet = openpyxl.load_workbook(
                Path(directory, "groups_available.xlsx")
            )["Subset Stats"]

            self.assertIsNotNone(find_row(worksheet, "Group OnlyGroup"))

    def test_group_mode_deduplicates_overlapping_weighted_memberships(self):
        viewer = self.make_viewer(
            ["Y", "Y", "A"],
            None,
            [{"G1"}, {"G1", "G2"}, set()],
        )
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(
                directory,
                [
                    "groups", "gmax", "10%", "cmin", "100%",
                    "id", "100%", "overlap",
                ],
                viewer=viewer,
                identity_weights=[0.5, 0.25, 1.0],
            )
            worksheet = openpyxl.load_workbook(
                Path(directory, "overlap.xlsx")
            )["Subset Stats"]
            self.assertEqual(
                worksheet.cell(find_row(worksheet, "Group G1"), 12).value,
                "Y1",
            )
            self.assertEqual(
                worksheet.cell(find_row(worksheet, "Group G2"), 12).value,
                "Y1",
            )

    def test_different_amino_acids_form_independent_exclusion_pools(self):
        viewer = self.make_viewer(
            ["Y", "Y", "F", "F", "Y", "A"],
            None,
            [{"GY"}, {"GY"}, {"GF"}, {"GF"}, set(), set()],
        )
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(
                directory,
                ["groups", "gmax", "40%", "cmin", "100%", "independent"],
                viewer=viewer,
            )
            worksheet = openpyxl.load_workbook(
                Path(directory, "independent.xlsx")
            )["Subset Stats"]
            self.assertEqual(
                worksheet.cell(find_row(worksheet, "Group GY"), 10).value,
                "Y1",
            )
            self.assertEqual(
                worksheet.cell(find_row(worksheet, "Group GF"), 10).value,
                "F1",
            )

    def test_gmax_remains_a_strict_upper_bound(self):
        # Excluding group G (rows 0 and 1) leaves Y, A outside: an outside Y
        # frequency of exactly 50%. gmax is exclusive, so 50% rejects Y1.
        self.assertEqual(
            label._calculate_outside_frequency(
                "Y",
                global_counts={"Y": 3},
                global_size=4,
                excluded_count=2,
                excluded_size=2,
            ),
            0.5,
        )
        for gmax, expected in (("50%", None), ("51%", "Y1")):
            viewer = self.make_viewer(
                ["Y", "Y", "Y", "A"],
                None,
                [{"G"}, {"G"}, set(), set()],
            )
            with self.subTest(gmax=gmax), tempfile.TemporaryDirectory() as directory:
                self.run_label(
                    directory,
                    ["groups", "gmax", gmax, "cmin", "100%", "boundary"],
                    viewer=viewer,
                )
                worksheet = openpyxl.load_workbook(
                    Path(directory, "boundary.xlsx")
                )["Subset Stats"]
                self.assertEqual(
                    worksheet.cell(find_row(worksheet, "Group G"), 10).value,
                    expected,
                )

    def test_subset_below_cmin_is_not_excluded_or_reported(self):
        viewer = self.make_viewer(
            ["Y", "Y", "Y", "A", "A", "A"],
            None,
            [{"High"}, {"High"}, {"Low"}, {"Low"}, set(), set()],
        )
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(
                directory,
                ["groups", "gmax", "20%", "cmin", "75%", "below_cmin"],
                viewer=viewer,
            )
            worksheet = openpyxl.load_workbook(
                Path(directory, "below_cmin.xlsx")
            )["Subset Stats"]
            header_row = find_row(worksheet, "Subset Name")
            self.assertIsNone(worksheet.cell(header_row, 10).value)

    def test_empty_shared_background_is_not_reported(self):
        viewer = self.make_viewer(
            ["Y", "Y"],
            None,
            [{"G1"}, {"G2"}],
        )
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(
                directory,
                ["groups", "cmin", "100%", "empty_background"],
                viewer=viewer,
            )
            worksheet = openpyxl.load_workbook(
                Path(directory, "empty_background.xlsx")
            )["Subset Stats"]
            header_row = find_row(worksheet, "Subset Name")
            self.assertIsNone(worksheet.cell(header_row, 10).value)

    def test_multiple_passing_amino_acids_share_one_ordered_cell(self):
        viewer = self.make_viewer(
            ["Y", "Y", "F", "A", "A"],
            None,
            [{"Mixed"}, {"Mixed"}, {"Mixed"}, set(), set()],
        )
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(
                directory,
                ["groups", "gmax", "40%", "cmin", "30%", "multi_residue"],
                viewer=viewer,
            )
            worksheet = openpyxl.load_workbook(
                Path(directory, "multi_residue.xlsx")
            )["Subset Stats"]
            self.assertEqual(
                worksheet.cell(find_row(worksheet, "Group Mixed"), 10).value,
                "Y1 | F1",
            )

    def label_sheets(self, sequences, groups, args, name, **kwargs):
        """Run label on a one-column alignment; return its two sheets and workbook."""
        viewer = self.make_viewer(sequences, None, groups)
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(directory, args + [name], viewer=viewer, **kwargs)
            workbook = openpyxl.load_workbook(Path(directory, f"{name}.xlsx"))
        return workbook["Subset Stats"], workbook["Occupancy Stats"]

    def test_ambiguity_codes_are_occupied_but_never_reported_as_conserved(self):
        # Rows 0-2 are the group: two of the residue and a gap, so 2/3 of the
        # group is the residue, enough for a cmin of 60%, and nothing else
        # holds it outside. As Y the cell reads "Y1"; as an ambiguity code it
        # is left empty, but the position is as occupied as with Y.
        groups = [{"G"}, {"G"}, {"G"}, set(), set()]
        args = ["groups", "gmax", "50%", "cmin", "60%"]
        control_subset, control_occupancy = self.label_sheets(
            ["Y", "Y", "-", "A", "C"], groups, args, "control"
        )
        self.assertEqual(control_subset.cell(find_row(control_subset, "Group G"), 10).value, "Y1")
        control_fill = control_occupancy.cell(
            find_row(control_occupancy, "Group G"), 10
        ).fill.fgColor.rgb
        for symbol in ("X", "B", "Z", "J", "*", "~", "_"):
            with self.subTest(symbol=symbol):
                subset, occupancy = self.label_sheets(
                    [symbol, symbol, "-", "A", "C"], groups, args, "ambiguous"
                )
                header_row = find_row(subset, "Subset Name")
                self.assertIsNone(subset.cell(header_row, 10).value)
                self.assertIsNone(subset.cell(find_row(subset, "Group G"), 10).value)
                # The position stays in the occupancy sheet, shaded as for Y.
                self.assertEqual(occupancy.cell(find_row(occupancy, "Subset Name"), 10).value, "#1")
                self.assertEqual(
                    occupancy.cell(find_row(occupancy, "Group G"), 10).fill.fgColor.rgb,
                    control_fill,
                )

    def test_ambiguity_codes_dilute_frequencies_like_any_residue(self):
        # Five rows in the group: 60% A and 40% X. A is reported at cmin 60%
        # (and below), X is never reported, and 61% is out of A's reach.
        sequences = ["A", "A", "A", "X", "X", "C", "C"]
        groups = [{"G"}] * 5 + [set(), set()]
        for cmin, expected in (("30%", "A1"), ("60%", "A1"), ("61%", None)):
            with self.subTest(cmin=cmin):
                subset, _occupancy = self.label_sheets(
                    sequences, groups, ["groups", "gmax", "50%", "cmin", cmin], "mixed"
                )
                self.assertEqual(
                    subset.cell(find_row(subset, "Group G"), 10).value, expected
                )

    def test_ambiguity_codes_are_not_reported_with_identity_weights_either(self):
        # Weights 1, 1, .5, .5, .5 over the group make A 2.5 of 3.5 (71.4%)
        # and X 1 of 3.5 (28.6%); X would pass a cmin of 20% unweighted.
        sequences = ["A", "A", "A", "X", "X", "C", "C"]
        groups = [{"G"}] * 5 + [set(), set()]
        weights = [1.0, 1.0, 0.5, 0.5, 0.5, 1.0, 1.0]
        for cmin, expected in (("20%", "A1"), ("70%", "A1"), ("72%", None)):
            with self.subTest(cmin=cmin):
                viewer = self.make_viewer(sequences, None, groups)
                with tempfile.TemporaryDirectory() as directory:
                    self.run_label(
                        directory,
                        ["groups", "id", "100%", "gmax", "50%", "cmin", cmin, "weighted"],
                        viewer=viewer,
                        identity_weights=weights,
                    )
                    worksheet = openpyxl.load_workbook(
                        Path(directory, "weighted.xlsx")
                    )["Subset Stats"]
                viewer.identity_weight_mock.assert_called_once()
                self.assertEqual(
                    worksheet.cell(find_row(worksheet, "Group G"), 12).value, expected
                )

    def test_global_conserved_row_never_lists_an_ambiguity_code(self):
        # 33 of 34 rows is 97.06%: conserved for U and O, which are residues,
        # but not for the ambiguity codes.
        cases = (
            ("X", ["None"]), ("B", ["None"]), ("Z", ["None"]), ("J", ["None"]),
            ("U", ["U1"]), ("O", ["O1"]),
        )
        for residue, expected in cases:
            sequences = [residue] * 33 + ["A"]
            viewer = self.make_viewer(
                sequences, [0] * len(sequences), [set() for _ in sequences]
            )
            with self.subTest(residue=residue), tempfile.TemporaryDirectory() as directory:
                self.run_label(directory, ["global_conserved"], viewer=viewer)
                workbook = openpyxl.load_workbook(Path(directory, "global_conserved.xlsx"))
                for sheet_name in ("Subset Stats", "Occupancy Stats"):
                    worksheet = workbook[sheet_name]
                    row = find_row(worksheet, "Global Conserved (>97%)")
                    values = [
                        cell.value for cell in worksheet[row + 1] if cell.value is not None
                    ]
                    self.assertEqual(values, expected, sheet_name)

    def test_selenocysteine_and_pyrrolysine_remain_conserved_candidates(self):
        groups = [{"G"}, {"G"}, set(), set()]
        for residue in ("U", "O"):
            with self.subTest(residue=residue):
                subset, _occupancy = self.label_sheets(
                    [residue, residue, "A", "C"],
                    groups,
                    ["groups", "gmax", "50%", "cmin", "100%"],
                    "rare",
                )
                self.assertEqual(
                    subset.cell(find_row(subset, "Group G"), 10).value, f"{residue}1"
                )

    def test_filename_cannot_escape_output_directory(self):
        with self.assertRaisesRegex(ValueError, "path separators"):
            label._normalize_output_filename("../outside")

    def test_enqueued_label_keeps_invocation_time_alignment_and_memberships(self):
        with tempfile.TemporaryDirectory() as directory:
            alignment = AlignmentStub()
            alignment.viewer_to_aln = np.array([0, 1, 2, -1])
            scheduler = CapturingScheduler()
            viewer = SimpleNamespace(
                alignment=alignment,
                active_reference="node0",
                alignment_offset=7,
                full_headers=["node0", "node1", "node2", "unaligned"],
                n_nodes=4,
                cluster_labels=np.array([0, 0, 1, -1]),
                group_labels=[{"GroupA"}, set(), set(), set()],
                console_text=SimpleNamespace(text=""),
                last_cluster_params=("leiden_1.0", 10),
                background_job_scheduler=scheduler,
            )
            with mock.patch.object(label, "CLUSTER_LABEL_DIRECTORY", directory), \
                    mock.patch.object(label.cfg, "INPUT_HDF5", "network.h5"):
                label.run(viewer, ["snapshot_report"])
                explicit_job = scheduler.job
                label.run(viewer, [])
                automatic_job = scheduler.job

            # Rows are materialized from the sparse matrix, so mutate the matrix.
            alignment.aln.matrix[0, 0] = Alignment_Manager.AA_TO_INT["G"]
            self.assertEqual(str(alignment.aln[0].seq), "G")
            alignment.col_to_label[0] = "99"
            viewer.cluster_labels[0] = 9
            viewer.group_labels[0].add("LaterGroup")
            viewer.active_reference = "node2"
            viewer.alignment_offset = 99

            snapshot = explicit_job["payload"].viewer_snapshot
            self.assertEqual(str(snapshot.alignment.aln[0].seq), "A")
            self.assertEqual(snapshot.alignment.col_to_label, {0: "1"})
            self.assertEqual(snapshot.cluster_labels, (0, 0, 1, -1))
            self.assertEqual(snapshot.group_labels[0], frozenset({"GroupA"}))
            self.assertEqual(snapshot.active_reference, "node0")
            self.assertEqual(snapshot.alignment_offset, 7)
            self.assertTrue(explicit_job["allow_overwrite"])
            self.assertTrue(snapshot._label_allow_overwrite)
            self.assertFalse(automatic_job["allow_overwrite"])
            self.assertFalse(
                automatic_job["payload"].viewer_snapshot._label_allow_overwrite
            )

    def reported_cell(self, directory, filename, subset, column=10):
        worksheet = openpyxl.load_workbook(
            Path(directory, filename)
        )["Subset Stats"]
        return worksheet.cell(find_row(worksheet, subset), column).value

    def test_cmin_allows_for_float_noise_in_weighted_frequencies(self):
        # Eight weights of 0.1 add up to 0.7999999999999999 one at a time but
        # to 0.8 pairwise, so a residue in every row scored just under 100%.
        viewer = self.make_viewer(
            ["Y"] * 8 + ["A"] * 2, None, [{"G"}] * 8 + [set()] * 2
        )
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(
                directory,
                ["groups", "gmax", "50%", "cmin", "100%", "id", "100%", "noisy"],
                viewer=viewer,
                identity_weights=[0.1] * 8 + [1.0, 1.0],
            )
            self.assertEqual(self.reported_cell(directory, "noisy.xlsx", "Group G", 12), "Y1")

    def test_cmin_99_9_percent_accepts_a_residue_in_999_of_1000_rows(self):
        viewer = self.make_viewer(
            ["Y"] * 999 + ["A"] * 11, None, [{"G"}] * 1000 + [set()] * 10
        )
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(
                directory,
                ["groups", "gmax", "50%", "cmin", "99.9%", "thousand"],
                viewer=viewer,
            )
            self.assertEqual(self.reported_cell(directory, "thousand.xlsx", "Group G"), "Y1")

    def test_gmax_rejects_a_value_equal_to_it_despite_float_noise(self):
        # Outside G the weighted Y frequency is 0.1 / 0.2 = 50%, which the
        # sums give as 0.4999999999999999.
        for gmax, expected in (("50%", None), ("51%", "Y1")):
            viewer = self.make_viewer(["Y", "Y", "A"], None, [{"G"}, set(), set()])
            with self.subTest(gmax=gmax), tempfile.TemporaryDirectory() as directory:
                self.run_label(
                    directory,
                    ["groups", "gmax", gmax, "cmin", "99%", "id", "100%", "noisy"],
                    viewer=viewer,
                    identity_weights=[0.1, 0.1, 0.1],
                )
                self.assertEqual(
                    self.reported_cell(directory, "noisy.xlsx", "Group G", 12), expected
                )

    def test_label_without_a_defined_cluster_or_group_fails_before_queueing(self):
        # The viewer holds an all -1 cluster array and a list of empty group
        # sets, not None, when nothing is defined.
        cases = (
            ([], [-1, -1, -1], "No clusters or groups defined."),
            (["clusters"], [-1, -1, -1], "Run 'cluster' first."),
            (["groups"], [-1, -1, -1], "No groups defined."),
            (["groups"], [0, 0, 1], "No groups defined."),
        )
        for args, cluster_labels, expected in cases:
            viewer = self.make_viewer(
                ["A", "A", "C"], cluster_labels, [set(), set(), set()]
            )
            viewer.background_job_scheduler = CapturingScheduler()
            with self.subTest(args=args, cluster_labels=cluster_labels), \
                    mock.patch.object(label.Command_Engine, "command_failed") as failed, \
                    redirect_stdout(io.StringIO()):
                label.run(viewer, args)

            failed.assert_called_once()
            self.assertEqual(str(failed.call_args.args[1]), "Error: " + expected)
            self.assertEqual(viewer.console_text.text, "Error: " + expected)
            self.assertIsNone(viewer.background_job_scheduler.job)

    def test_label_queues_when_only_one_kind_of_subset_is_defined(self):
        for args, cluster_labels, group_labels in (
            ([], [-1, -1, -1], [{"G"}, set(), set()]),
            ([], [0, 0, -1], [set(), set(), set()]),
            (["groups"], [-1, -1, -1], [{"G"}, set(), set()]),
            (["clusters"], [0, 0, -1], [set(), set(), set()]),
        ):
            viewer = self.make_viewer(["A", "A", "C"], cluster_labels, group_labels)
            viewer.background_job_scheduler = CapturingScheduler()
            with self.subTest(args=args), tempfile.TemporaryDirectory() as directory, \
                    mock.patch.object(label, "CLUSTER_LABEL_DIRECTORY", directory), \
                    mock.patch.object(
                        label.Command_Engine,
                        "get_alignment_mapping",
                        return_value=(np.arange(3), np.arange(3)),
                    ), redirect_stdout(io.StringIO()):
                label.run(viewer, args)

            self.assertIsNotNone(viewer.background_job_scheduler.job)

    def test_double_dash_help_is_help_and_writes_no_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            with redirect_stdout(io.StringIO()) as output:
                self.run_label(directory, ["--help"])

            self.assertIn("Differential Labeling", output.getvalue())
            self.assertEqual(list(Path(directory).iterdir()), [])

    def test_help_needs_no_alignment_or_reference(self):
        no_reference = AlignmentStub()
        no_reference.has_reference = False
        for alignment in (None, no_reference):
            for token in ("help", "-h", "-?", "--help"):
                viewer = SimpleNamespace(
                    alignment=alignment, console_text=SimpleNamespace(text="")
                )
                with self.subTest(alignment=alignment, token=token), \
                        mock.patch.object(label.Command_Engine, "command_failed") as failed, \
                        redirect_stdout(io.StringIO()) as output:
                    label.run(viewer, [token])

                failed.assert_not_called()
                self.assertIn("Differential Labeling", output.getvalue())
                self.assertEqual(
                    viewer.console_text.text, "Help information printed to the Console"
                )

    def test_help_describes_reset_and_the_exclusive_gmax(self):
        with redirect_stdout(io.StringIO()) as output:
            label.print_help()

        self.assertIn("label reset", output.getvalue())
        self.assertIn("is rejected when its outside frequency is >= gmax", " ".join(
            output.getvalue().split()
        ))

    def test_mistyped_assignments_are_rejected_not_taken_as_filenames(self):
        for args in (
            ["gmax=0.4"], ["cmin=90%"], ["id=90"], ["-x"], ["0.4", "--cmin"],
            ["cmin", "90%", "-h"],
        ):
            with self.subTest(args=args), tempfile.TemporaryDirectory() as directory:
                viewer = self.run_label(directory, args)

                self.assertIn("Unrecognized argument", viewer.console_text.text)
                self.assertEqual(list(Path(directory).iterdir()), [])
                viewer.identity_weight_mock.assert_not_called()

    def test_thresholds_outside_zero_to_one_hundred_percent_are_rejected(self):
        for args in (
            ["cmin", "-1"], ["gmax", "-5"], ["cmin", "150%"], ["gmax", "101"],
            ["cmin", "-1%"], ["-5"], ["0.4", "150%"],
        ):
            with self.subTest(args=args), tempfile.TemporaryDirectory() as directory:
                viewer = self.run_label(directory, args)

                self.assertIn("outside the supported range", viewer.console_text.text)
                self.assertEqual(list(Path(directory).iterdir()), [])

    def test_thresholds_of_zero_and_one_hundred_percent_are_accepted(self):
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(directory, ["gmax", "0", "cmin", "100%", "bounds"])

            self.assertTrue(Path(directory, "bounds.xlsx").is_file())

    def test_label_parameters_show_percentages_without_truncation(self):
        # int(0.29 * 100) is 28 and int(0.57 * 100) is 56.
        for gmax, cmin, expected in (
            ("57%", "29%", "gmax_outside=57%, cmin=29%, target=all"),
            ("58%", "0.5%", "gmax_outside=58%, cmin=0.5%, target=all"),
        ):
            with self.subTest(gmax=gmax, cmin=cmin), \
                    tempfile.TemporaryDirectory() as directory:
                self.run_label(directory, ["gmax", gmax, "cmin", cmin, "params"])

                metadata = openpyxl.load_workbook(
                    Path(directory, "params.xlsx")
                )["Meta Data"]
                self.assertEqual(
                    metadata.cell(find_row(metadata, "Label Parameters"), 2).value,
                    expected,
                )

    def test_hex_colors_follow_the_viewer_when_a_cluster_has_no_aligned_node(self):
        # Cluster 1's nodes are not in the alignment, but cluster.run colours
        # the viewer from every cluster id, so the report must as well.
        def viewer_with_clusters():
            return SimpleNamespace(
                alignment=AlignmentStub(),
                active_reference="node0",
                full_headers=["node0", "node1", "node2", "other0", "other1"],
                n_nodes=5,
                cluster_labels=np.array([0, 0, 2, 1, 1]),
                group_labels=[set() for _ in range(5)],
                console_text=SimpleNamespace(text=""),
                last_cluster_params=None,
            )

        viewer_colors = label.cluster_cmd.get_cluster_color_map([0, 1, 2])
        self.assertNotEqual(
            viewer_colors[2], label.cluster_cmd.get_cluster_color_map([0, 2])[2]
        )
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(
                directory, ["clusters", "colors"], viewer=viewer_with_clusters(),
                real_colors=True,
            )

            for cluster_id in (0, 2):
                self.assertEqual(
                    self.reported_cell(
                        directory, "colors.xlsx", f"Cluster {cluster_id}", 4
                    ),
                    label.mcolors.to_hex(viewer_colors[cluster_id]),
                )

    def test_subsets_without_an_aligned_node_are_named_in_a_terminal_warning(self):
        viewer = SimpleNamespace(
            alignment=AlignmentStub(),
            active_reference="node0",
            full_headers=["node0", "node1", "node2", "other0", "other1"],
            n_nodes=5,
            cluster_labels=np.array([0, 0, 2, 1, 1]),
            group_labels=[{"Seen"}, set(), set(), {"Only"}, set()],
            console_text=SimpleNamespace(text=""),
            last_cluster_params=None,
        )
        with tempfile.TemporaryDirectory() as directory:
            with redirect_stdout(io.StringIO()) as output:
                self.run_label(directory, ["warned"], viewer=viewer)

            self.assertIn(
                "Warning: no aligned sequence for Cluster 1, Group Only; "
                "not included in the report.",
                output.getvalue(),
            )
            worksheet = openpyxl.load_workbook(
                Path(directory, "warned.xlsx")
            )["Subset Stats"]
            names = {
                worksheet.cell(row, 1).value for row in range(1, worksheet.max_row + 1)
            }
            self.assertLessEqual({"Cluster 0", "Cluster 2", "Group Seen"}, names)
            self.assertNotIn("Cluster 1", names)
            self.assertNotIn("Group Only", names)

    def test_no_warning_when_every_subset_has_an_aligned_node(self):
        with tempfile.TemporaryDirectory() as directory:
            with redirect_stdout(io.StringIO()) as output:
                self.run_label(directory, [])

            self.assertNotIn("Warning", output.getvalue())

    @unittest.skipUnless(os.name == "nt", "Only Windows refuses to replace an open file.")
    def test_report_held_open_elsewhere_fails_before_the_analysis(self):
        import ctypes
        from ctypes import wintypes

        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel32.CreateFileW.restype = wintypes.HANDLE
        kernel32.CreateFileW.argtypes = [
            wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
            wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE,
        ]
        kernel32.CloseHandle.argtypes = [wintypes.HANDLE]
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory, "open_report.xlsx")
            output_path.write_bytes(b"open in Excel")
            # Excel reads a workbook with other readers allowed but no writer.
            handle = kernel32.CreateFileW(
                str(output_path), 0x80000000, 0x1, None, 3, 0x80, None
            )
            self.assertNotEqual(handle, wintypes.HANDLE(-1).value)
            try:
                viewer = self.make_viewer(["A", "A", "C"], [0, 0, 1], [set()] * 3)
                viewer.background_job_scheduler = CapturingScheduler()
                with mock.patch.object(label, "CLUSTER_LABEL_DIRECTORY", directory), \
                        redirect_stdout(io.StringIO()):
                    label.run(viewer, ["open_report"])
            finally:
                kernel32.CloseHandle(handle)

            self.assertIsNone(viewer.background_job_scheduler.job)
            self.assertIn("cannot be replaced", viewer.console_text.text)
            self.assertIn(str(output_path), viewer.console_text.text)
            self.assertEqual(output_path.read_bytes(), b"open in Excel")

    def test_replaceable_existing_report_is_still_queued(self):
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory, "again.xlsx")
            output_path.write_bytes(b"previous")
            viewer = self.make_viewer(["A", "A", "C"], [0, 0, 1], [set()] * 3)
            viewer.background_job_scheduler = CapturingScheduler()
            with mock.patch.object(label, "CLUSTER_LABEL_DIRECTORY", directory), \
                    mock.patch.object(
                        label.Command_Engine,
                        "get_alignment_mapping",
                        return_value=(np.arange(3), np.arange(3)),
                    ), redirect_stdout(io.StringIO()):
                label.run(viewer, ["again"])

            self.assertIsNotNone(viewer.background_job_scheduler.job)
            self.assertEqual(output_path.read_bytes(), b"previous")

    def test_underscored_numbers_and_trailing_dots_in_filenames(self):
        for name, expected in (
            ("2026_01_01", "2026_01_01.xlsx"),
            ("report.xlsx.", "report.xlsx"),
            ("report.", "report.xlsx"),
        ):
            with self.subTest(name=name), tempfile.TemporaryDirectory() as directory:
                self.run_label(directory, [name])

                self.assertEqual(
                    [path.name for path in Path(directory).iterdir()], [expected]
                )


class StatusLine:
    """Stands in for the Viewer the scheduler reports its status to."""

    def set_background_job_status(self, message):
        pass


class LabelJobCancellationTests(unittest.TestCase):
    """A label job that is told to stop while it runs ends as cancelled and leaves no file.

    The job runs on the real background job scheduler's thread, and each test asks
    for the stop from inside the analysis or the workbook's save, at the point named.
    """

    @classmethod
    def setUpClass(cls):
        from PySide6 import QtCore

        cls.qt = QtCore
        cls.application = QtCore.QCoreApplication.instance() or QtCore.QCoreApplication([])

    def setUp(self):
        from Background_Job_Scheduler import BackgroundJobScheduler

        self.runner = LabelWorkbookPercentTests()
        self.status = StatusLine()
        self.scheduler = BackgroundJobScheduler(self.status)
        self.addCleanup(self.scheduler.shutdown)
        # A finished label job asks for its folder to be shown, to a stand-in that has windows.
        self.reveal = self.enterContext(mock.patch("Background_Job_Scheduler.open_in_file_manager"))
        self.reports = []
        self.scheduler.job_succeeded.connect(lambda job, result, elapsed: self.reports.append("succeeded"))
        self.scheduler.job_failed.connect(lambda job, error, detail, elapsed: self.reports.append("failed"))
        self.scheduler.job_cancelled.connect(lambda job, started: self.reports.append("stopped" if started else "removed"))
        self.stack = self.enterContext(redirect_stdout(io.StringIO()))

    def settle(self):
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            self.application.processEvents()
            if self.scheduler.queue_depth == 0 and self.reports:
                break
            time.sleep(0.005)
        self.application.processEvents()

    def run_label(self, directory, args):
        return self.runner.run_label(directory, args, scheduler=self.scheduler, after_run=lambda viewer: self.settle())

    def stop_job(self):
        """Tell the running label job to stop, as the Session card's remove does."""
        (running,) = [entry for entry in self.scheduler.snapshot() if entry.running]
        self.assertTrue(self.scheduler.cancel(running.job.job_id))

    def stop_when_called(self, owner, name, then=None):
        """Patch owner.name to ask for the stop on its first call, and then do what it did."""
        real = getattr(owner, name)
        calls = []

        def stopping(*args, **kwargs):
            if not calls:
                self.stop_job()
                if then is not None:
                    then()
            calls.append(name)
            return real(*args, **kwargs)

        return mock.patch.object(owner, name, side_effect=stopping), calls

    def assert_stopped_leaving_nothing(self, directory):
        self.assertEqual(self.reports, ["stopped"])
        self.assertEqual(os.listdir(directory), [])
        self.assertEqual(self.scheduler.snapshot(), [])

    def test_a_stop_during_the_analysis_ends_the_job_with_no_file(self):
        # Each stage is a call the analysis makes; the stop comes during it, and the next check point ends the job.
        for owner, name in (
            (label, "get_sequence_stats"),
            (label, "_calculate_weighted_frequencies"),
            (label, "_build_label_tasks"),
            (label, "_calculate_outside_frequency"),
        ):
            with self.subTest(stage=name), tempfile.TemporaryDirectory() as directory:
                self.reports.clear()
                patch, calls = self.stop_when_called(owner, name)
                with patch, mock.patch.object(label.tempfile, "mkstemp", wraps=label.tempfile.mkstemp) as made:
                    self.run_label(directory, [])
                self.assertTrue(calls, "the stage ran")
                self.assert_stopped_leaving_nothing(directory)
                # The job stopped at its next check point, before it built the workbook or made its partial file.
                made.assert_not_called()

    def test_a_stop_while_the_workbook_is_saved_deletes_the_partial_file(self):
        partial_seen = []

        class StopsOnSecondWrite(label.CancellableFile):
            def write(inner, data):
                writes.append(1)
                if len(writes) == 2:
                    partial_seen.extend(path.name for path in Path(directory).iterdir())
                    self.stop_job()
                return super().write(data)

        writes = []
        with tempfile.TemporaryDirectory() as directory:
            with mock.patch.object(label, "CancellableFile", StopsOnSecondWrite):
                self.run_label(directory, ["report"])
            self.assertEqual(len(partial_seen), 1)
            self.assertIn(".partial.xlsx", partial_seen[0], "the partial file existed when the stop came")
            self.assert_stopped_leaving_nothing(directory)

    def test_a_stop_leaves_an_existing_report_of_that_name_as_it_was(self):
        with tempfile.TemporaryDirectory() as directory:
            existing = Path(directory, "report.xlsx")
            existing.write_bytes(b"previous content")
            patch, calls = self.stop_when_called(label, "_calculate_weighted_frequencies")
            with patch:
                self.run_label(directory, ["report"])
            self.assertEqual(self.reports, ["stopped"])
            self.assertEqual(os.listdir(directory), ["report.xlsx"])
            self.assertEqual(existing.read_bytes(), b"previous content")

    def test_the_output_is_free_again_and_the_next_job_runs_after_a_stop(self):
        with tempfile.TemporaryDirectory() as directory:
            patch, _calls = self.stop_when_called(label, "_calculate_weighted_frequencies")
            with patch:
                self.run_label(directory, ["report"])
            self.assertFalse(self.scheduler.is_output_path_reserved(os.path.join(directory, "report.xlsx")))
            self.reports.clear()
            self.run_label(directory, ["report"])
            self.assertEqual(self.reports, ["succeeded"])
            self.assertEqual(os.listdir(directory), ["report.xlsx"])
            self.assertIn("Meta Data", openpyxl.load_workbook(Path(directory, "report.xlsx")).sheetnames)

    def close_the_viewer_during_the_save(self, directory, *, job_obeys, wait=2.5):
        """Run a label job and close the scheduler, as closing the Viewer does, at the first write of its workbook.

        With ``job_obeys`` the job waits at that write until it is told to stop, as a job
        does that checks; without, it holds there until the test lets it go, as a job stuck
        in one long call does. Returns (how long closing took, what the folder held when
        closing returned, which is before the job is let go).
        """
        entered, release, test = threading.Event(), threading.Event(), self
        took = []

        class HoldsAtFirstWrite(label.CancellableFile):
            held = False

            def write(inner, data):
                if data and not HoldsAtFirstWrite.held:
                    HoldsAtFirstWrite.held = True
                    entered.set()
                    if job_obeys:
                        while not cancel_requested():
                            time.sleep(0.005)
                    else:
                        release.wait(30)
                return super().write(data)

        def close(viewer):
            self.assertTrue(entered.wait(30))
            started = time.monotonic()
            self.scheduler.shutdown(wait=wait)
            took.append((time.monotonic() - started, sorted(os.listdir(directory))))
            release.set()
            deadline = time.monotonic() + 30
            while self.scheduler._thread.is_alive() and time.monotonic() < deadline:
                time.sleep(0.005)

        with mock.patch.object(label, "CancellableFile", HoldsAtFirstWrite):
            self.runner.run_label(directory, ["report"], scheduler=self.scheduler, after_run=close)
        return took[0]

    def test_closing_the_viewer_during_the_save_leaves_neither_a_partial_file_nor_the_report(self):
        for job_obeys in (True, False):
            with self.subTest(job_obeys=job_obeys), tempfile.TemporaryDirectory() as directory:
                took, left_when_closed = self.close_the_viewer_during_the_save(directory, job_obeys=job_obeys, wait=0.4)
                self.assertLess(took, 3.0, "closing returns within the bound")
                # Before the job that ignores the stop is let go: the scheduler deleted the partial file itself.
                self.assertEqual(left_when_closed, [])
                self.assertEqual(os.listdir(directory), [])
                self.scheduler = type(self.scheduler)(self.status)   # the next subtest's scheduler
                self.addCleanup(self.scheduler.shutdown)

    def test_closing_the_viewer_during_a_report_leaves_an_earlier_finished_one_as_it_was(self):
        with tempfile.TemporaryDirectory() as directory:
            self.run_label(directory, ["first"])
            self.assertEqual(self.reports, ["succeeded"])
            before = Path(directory, "first.xlsx").read_bytes()
            self.close_the_viewer_during_the_save(directory, job_obeys=False, wait=0.3)
            self.assertEqual(os.listdir(directory), ["first.xlsx"])
            self.assertEqual(Path(directory, "first.xlsx").read_bytes(), before)

    def test_a_stop_is_a_cancellation_not_a_failure_in_the_labels_own_error_handling(self):
        # label reports its failures as failed commands; a stop must not reach that handler.
        with tempfile.TemporaryDirectory() as directory:
            patch, _calls = self.stop_when_called(label, "_calculate_weighted_frequencies")
            with patch, mock.patch.object(label.Command_Engine, "command_failed") as failed:
                self.run_label(directory, [])
            failed.assert_not_called()
            self.assertEqual(self.reports, ["stopped"])


if __name__ == "__main__":
    unittest.main()
