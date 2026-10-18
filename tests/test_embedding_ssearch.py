"""Embedding_SSEARCH search results: scores, hit ranking and reports.

Covers finish_search (raw score, alignment length and percent identity along
the selected path), score normalization, the ranked-hit filter and the text
and metadata-viewer workbook reports. Hardware plan selection is covered in
test_embedding_ssearch_hardware.
"""

import io
import os
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from io import StringIO
from types import SimpleNamespace
from unittest import mock

import numpy as np
import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_DIR = os.path.join(PROJECT_ROOT, "src")
TOOLS_DIR = os.path.join(SRC_DIR, "tools")
for directory in (SRC_DIR, TOOLS_DIR):
    if directory not in sys.path:
        sys.path.insert(0, directory)

import Command_Engine  # noqa: E402
from web_ui import meta_backend  # noqa: E402

with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
    import Embedding_SSEARCH


class FinishSearchTests(unittest.TestCase):
    def test_ssearch_uses_only_the_selected_mode_and_preserves_matrix(self):
        matrix = np.array(
            [
                [3.0, -1.0],
                [-1.0, 3.0],
            ],
            dtype=np.float32,
        )
        original = matrix.copy()

        local_result = Embedding_SSEARCH.finish_search(
            (0, "local", 2, 2, "local", -2.0, "longer_sequence", matrix,
             "AC", "AD")
        )
        np.testing.assert_array_equal(matrix, original)
        self.assertEqual(float(local_result["raw_score"]), 2.0)
        self.assertEqual(int(local_result["aln_len"]), 2)
        self.assertEqual(local_result["identity"], 50.0)

        global_result = Embedding_SSEARCH.finish_search(
            (1, "global", 2, 2, "global", 0.0, "alignment_length", matrix,
             "AC", "AD")
        )
        np.testing.assert_array_equal(matrix, original)
        self.assertEqual(float(global_result["raw_score"]), 6.0)
        self.assertEqual(int(global_result["aln_len"]), 2)
        self.assertEqual(global_result["identity"], 50.0)

    def test_ssearch_identity_denominators_follow_alignment_mode(self):
        # Only the C/C pair aligns well, so a global path pays two end gaps
        # while the local path is that single pair. Transposing the case
        # moves the gaps from the target side to the query side.
        end_gap_matrix = np.array([[-5.0], [3.0], [-5.0]], dtype=np.float32)
        for matrix, query, target in (
            (end_gap_matrix, "ACD", "C"),
            (end_gap_matrix.T.copy(), "C", "ACD"),
        ):
            with self.subTest(query=query, target=target):
                rows, cols = matrix.shape
                local_result = Embedding_SSEARCH.finish_search(
                    (0, "t", rows, cols, "local", -2.0, "longer_sequence",
                     matrix, query, target)
                )
                global_result = Embedding_SSEARCH.finish_search(
                    (0, "t", rows, cols, "global", 0.0, "alignment_length",
                     matrix, query, target)
                )
                self.assertEqual(int(local_result["aln_len"]), 1)
                self.assertEqual(local_result["identity"], 100.0)
                self.assertEqual(int(global_result["aln_len"]), 3)
                self.assertAlmostEqual(global_result["identity"], 100.0 / 3.0)

        # A/A and W/W flank one unpaired residue, so the local path keeps the
        # internal gap in its denominator on either side of the alignment.
        internal_gap_matrix = np.full((3, 2), -10.0, dtype=np.float32)
        internal_gap_matrix[0, 0] = 4.0
        internal_gap_matrix[2, 1] = 4.0
        for matrix, query, target in (
            (internal_gap_matrix, "ACW", "AW"),
            (internal_gap_matrix.T.copy(), "AW", "ACW"),
        ):
            with self.subTest(query=query, target=target):
                rows, cols = matrix.shape
                result = Embedding_SSEARCH.finish_search(
                    (0, "t", rows, cols, "local", -1.0, "longer_sequence",
                     matrix, query, target)
                )
                self.assertEqual(int(result["aln_len"]), 3)
                self.assertAlmostEqual(result["identity"], 200.0 / 3.0)

    def test_ssearch_identity_counts_only_standard_amino_acids(self):
        residues = "AXBZJUOC"
        matrix = np.where(
            np.eye(len(residues), dtype=bool), 4.0, -10.0
        ).astype(np.float32)
        result = Embedding_SSEARCH.finish_search(
            (0, "t", len(residues), len(residues), "local", -2.0,
             "longer_sequence", matrix, residues, residues)
        )
        self.assertEqual(int(result["aln_len"]), len(residues))
        self.assertEqual(result["identity"], 25.0)

    def test_ssearch_rejects_sequences_that_do_not_match_the_matrix(self):
        matrix = np.zeros((2, 3), dtype=np.float32)
        with self.assertRaisesRegex(ValueError, "has shape"):
            Embedding_SSEARCH.finish_search(
                (0, "t", 2, 3, "local", -2.0, "longer_sequence", matrix,
                 "AC", "ACDE")
            )

    def test_reported_length_is_the_normalization_denominator(self):
        # Only A/A scores, so the global path is that pair plus four gap
        # columns: five columns for a 2-residue query and a 4-residue target.
        matrix = np.full((2, 4), -5.0, dtype=np.float32)
        matrix[0, 1] = 4.0
        cases = {
            "alignment_length": 5,
            "shorter_sequence": 2,
            "longer_sequence": 4,
            "average_sequence": 3.0,
        }
        for norm_mode, denominator in cases.items():
            with self.subTest(norm_mode=norm_mode):
                result = Embedding_SSEARCH.finish_search(
                    (0, "t", 2, 4, "global", 0.0, norm_mode, matrix, "AC", "DAEF")
                )
                self.assertEqual(float(result["raw_score"]), 4.0)
                self.assertEqual(int(result["aln_len"]), 5)
                self.assertEqual(result["identity"], 20.0)
                self.assertEqual(result["length"], denominator)
                self.assertAlmostEqual(
                    float(result["norm_score"]), 4.0 / denominator
                )
        # An unknown mode is an error, not alignment-length normalization.
        with self.assertRaisesRegex(ValueError, "Unknown NORM_MODE 'unrecognized_mode'"):
            Embedding_SSEARCH.finish_search(
                (0, "t", 2, 4, "global", 0.0, "unrecognized_mode", matrix, "AC", "DAEF")
            )


class FinishSearchBatchTests(unittest.TestCase):
    """finish_search_batch gives, for a padded microbatch, exactly what finish_search gives per target."""

    RESIDUES = list("ACDEFGHIKLMNPQRSTVWYXBZ")

    def microbatch(self, seed, mode, gap, norm_mode, count=9, rows=17, dense=False):
        rng = np.random.default_rng(seed)
        lengths = [int(n) for n in rng.integers(1, 30, size=count)]
        padded = max(lengths)
        if dense:
            matrices = (rng.random((count, rows, padded)) * 3.0).astype(np.float32)
        else:
            matrices = (rng.integers(-3, 4, size=(count, rows, padded)) * 0.5).astype(np.float32)
        for offset, length in enumerate(lengths):
            matrices[offset, :, length:] = 0.0
        query = "".join(rng.choice(self.RESIDUES, size=rows))
        tasks = [
            (index, f"h{index}", f"h{index}", None, mode, gap, norm_mode, query,
             "".join(rng.choice(self.RESIDUES, size=length)))
            for index, length in enumerate(lengths)
        ]
        return tasks, lengths, matrices

    @staticmethod
    def per_target(tasks, lengths, matrices):
        return [
            Embedding_SSEARCH.finish_search((
                task[0], task[1], matrices.shape[1], length, task[4], task[5], task[6],
                matrices[offset, :, :length], task[7], task[8],
            ))
            for offset, (task, length) in enumerate(zip(tasks, lengths))
        ]

    def assert_same_rows(self, got, want):
        self.assertEqual(got, want)
        for row_got, row_want in zip(got, want):
            self.assertEqual(
                {key: type(value) for key, value in row_got.items()},
                {key: type(value) for key, value in row_want.items()},
            )

    def test_rows_equal_the_per_target_rows_for_every_mode_and_normalization(self):
        settings = [
            ("local", -2.0), ("local", 0.0), ("local", -0.7), ("global", 0.0), ("global", -0.5),
        ]
        for mode, gap in settings:
            for norm in ("alignment_length", "shorter_sequence", "longer_sequence", "average_sequence"):
                for dense in (False, True):
                    with self.subTest(mode=mode, gap=gap, norm=norm, dense=dense):
                        tasks, lengths, matrices = self.microbatch(5, mode, gap, norm, dense=dense)
                        self.assert_same_rows(
                            Embedding_SSEARCH.finish_search_batch(tasks, lengths, matrices),
                            self.per_target(tasks, lengths, matrices),
                        )

    def test_a_non_finite_matrix_is_refused(self):
        tasks, lengths, matrices = self.microbatch(1, "local", -2.0, "longer_sequence")
        matrices[2, 3, 1] = np.nan
        with self.assertRaisesRegex(FloatingPointError, "non-finite"):
            Embedding_SSEARCH.finish_search_batch(tasks, lengths, matrices)

    def test_a_matrix_that_does_not_fit_its_sequence_is_refused_like_a_single_target(self):
        tasks, lengths, matrices = self.microbatch(2, "global", 0.0, "longer_sequence")
        lengths = list(lengths)
        lengths[4] += 1
        with self.assertRaises(ValueError) as batch:
            Embedding_SSEARCH.finish_search_batch(tasks, lengths, matrices)
        with self.assertRaises(ValueError) as single:
            Embedding_SSEARCH.finish_search((
                4, tasks[4][1], matrices.shape[1], lengths[4], "global", 0.0, "longer_sequence",
                np.zeros((matrices.shape[1], lengths[4]), np.float32), tasks[4][7], tasks[4][8],
            ))
        self.assertEqual(str(batch.exception), str(single.exception))

    def test_searches_that_differ_run_one_at_a_time(self):
        tasks, lengths, matrices = self.microbatch(3, "local", -2.0, "longer_sequence")
        other_query = tasks[3][:7] + ("A" * matrices.shape[1], tasks[3][8])
        other_mode = tasks[5][:4] + ("global", 0.0, "alignment_length", tasks[5][7], tasks[5][8])
        mixed = list(tasks)
        mixed[3], mixed[5] = other_query, other_mode
        self.assert_same_rows(
            Embedding_SSEARCH.finish_search_batch(mixed, lengths, matrices),
            self.per_target(mixed, lengths, matrices),
        )

    def test_the_search_runs_the_float32_kernels_only_where_they_are_exact(self):
        from utilities import Network_Kernels

        cases = (
            ("global", 0.0, "_float32_global_length_identity_batch"),
            ("local", -2.0, "_float32_local_length_identity_batch"),
            ("local", 0.0, "_float32_local_length_identity_batch"),
        )
        for mode, gap, kernel in cases:
            tasks, lengths, matrices = self.microbatch(4, mode, gap, "longer_sequence")
            with self.subTest(mode=mode, gap=gap), mock.patch.object(
                Network_Kernels, kernel, wraps=getattr(Network_Kernels, kernel)
            ) as fast:
                # One target alone (the pool and scalar paths) and a microbatch both take them.
                Embedding_SSEARCH.finish_search(self.per_target_args(tasks, lengths, matrices, 0))
                self.assertEqual(fast.call_count, 1)
                Embedding_SSEARCH.finish_search_batch(tasks, lengths, matrices)
                self.assertEqual(fast.call_count, 2)
        # Other gaps, and matrices that are not float32, run the reference kernels.
        for mode, gap, scores in (("global", -0.5, np.float32), ("local", -0.7, np.float32),
                                  ("local", -2.0, np.float64), ("global", 0.0, np.float64)):
            tasks, lengths, matrices = self.microbatch(4, mode, gap, "longer_sequence")
            with self.subTest(mode=mode, gap=gap, dtype=scores.__name__), mock.patch.object(
                Network_Kernels, "_float32_global_length_identity_batch", side_effect=AssertionError("float32")
            ), mock.patch.object(
                Network_Kernels, "_float32_local_length_identity_batch", side_effect=AssertionError("float32")
            ):
                Embedding_SSEARCH.finish_search(self.per_target_args(tasks, lengths, matrices.astype(scores), 0))
                Embedding_SSEARCH.finish_search_batch(tasks, lengths, matrices.astype(scores))

    @staticmethod
    def per_target_args(tasks, lengths, matrices, offset):
        task = tasks[offset]
        return (task[0], task[1], matrices.shape[1], lengths[offset], task[4], task[5], task[6],
                matrices[offset, :, :lengths[offset]], task[7], task[8])

    def test_the_results_are_the_reference_kernels_bit_for_bit(self):
        from utilities import Network_Kernels

        for mode, gap in (("local", -2.0), ("global", 0.0)):
            tasks, lengths, matrices = self.microbatch(8, mode, gap, "longer_sequence", dense=True)
            rows = Embedding_SSEARCH.finish_search_batch(tasks, lengths, matrices)
            kernel = (Network_Kernels.local_score_length_identity if mode == "local"
                      else Network_Kernels.global_score_length_identity)
            for offset, (task, length, row) in enumerate(zip(tasks, lengths, rows)):
                raw, path_len, identities = kernel(
                    np.ascontiguousarray(matrices[offset, :, :length]), gap,
                    Embedding_SSEARCH.encode_identity_residues(task[7], Embedding_SSEARCH.QUERY_UNMATCHED_CODE),
                    Embedding_SSEARCH.encode_identity_residues(task[8], Embedding_SSEARCH.TARGET_UNMATCHED_CODE),
                )
                with self.subTest(mode=mode, target=offset):
                    self.assertEqual((row["raw_score"], row["aln_len"]), (raw, path_len))
                    self.assertEqual(row["identity"], 100.0 * int(identities) / path_len if path_len > 0 else 0.0)


class ScoreNormalizationTests(unittest.TestCase):
    def test_each_mode_divides_by_its_own_length(self):
        # Raw score 10 on a 4-column path between 5- and 8-residue sequences.
        cases = {
            "alignment_length": 2.5,
            "shorter_sequence": 2.0,
            "longer_sequence": 1.25,
            "average_sequence": 10.0 / 6.5,
        }
        for mode, expected in cases.items():
            with self.subTest(mode=mode):
                self.assertAlmostEqual(
                    Embedding_SSEARCH.normalize_score(10.0, 4, 5, 8, mode),
                    expected,
                )

    def test_unknown_mode_is_rejected(self):
        # Hand-written settings reach a direct CLI run unchecked (the Tools
        # window and MCP offer only the four modes), and an unknown mode used
        # to normalize by alignment length silently.
        for mode in ("unrecognized_mode", "Longer_Sequence", " longer_sequence", 5):
            with self.subTest(mode=mode):
                for alignment_mode in ("local", "global"):
                    with self.assertRaisesRegex(ValueError, "Unknown NORM_MODE"):
                        Embedding_SSEARCH.validate_score_normalization(alignment_mode, mode)
                with self.assertRaisesRegex(ValueError, "Unknown NORM_MODE"):
                    Embedding_SSEARCH.normalize_score(10.0, 4, 5, 8, mode)
        with mock.patch.object(Embedding_SSEARCH, "load_tool_settings"), mock.patch.multiple(
            Embedding_SSEARCH, ALIGNMENT_MODE="global", NORM_MODE="unrecognized_mode"
        ), mock.patch.object(Embedding_SSEARCH, "prepare_database_embeddings") as prepare:
            with self.assertRaisesRegex(ValueError, "Unknown NORM_MODE 'unrecognized_mode'"):
                Embedding_SSEARCH.main([])
        prepare.assert_not_called()

    def test_unknown_alignment_mode_is_rejected(self):
        # Only an exact "local" selected the local kernel and gap penalty, so a
        # hand-written "Local" or a typo ran a global search silently, and
        # "Local" with alignment_length also passed the local-score check.
        for alignment_mode in ("Local", "GLOBAL", " local", "glocal", "", None, 1):
            with self.subTest(alignment_mode=alignment_mode):
                for norm_mode in ("alignment_length", "longer_sequence"):
                    with self.assertRaisesRegex(ValueError, "Unknown ALIGNMENT_MODE"):
                        Embedding_SSEARCH.validate_score_normalization(alignment_mode, norm_mode)
        with mock.patch.object(Embedding_SSEARCH, "load_tool_settings"), mock.patch.multiple(
            Embedding_SSEARCH, ALIGNMENT_MODE="Local", NORM_MODE="longer_sequence"
        ), mock.patch.object(Embedding_SSEARCH, "prepare_database_embeddings") as prepare:
            with self.assertRaisesRegex(ValueError, "Unknown ALIGNMENT_MODE 'Local'"):
                Embedding_SSEARCH.main([])
        prepare.assert_not_called()

    def test_empty_lengths_normalize_to_zero(self):
        self.assertEqual(
            Embedding_SSEARCH.normalize_score(10.0, 0, 5, 8, "alignment_length"),
            0.0,
        )
        for mode in ("shorter_sequence", "longer_sequence", "average_sequence"):
            with self.subTest(mode=mode):
                self.assertEqual(
                    Embedding_SSEARCH.normalize_score(10.0, 4, 0, 0, mode), 0.0
                )


class RankedHitFilterTests(unittest.TestCase):
    @staticmethod
    def ranked(*, manual_query, top_k):
        # Rows arrive sorted by normalized score; the query's own record ranks
        # first because a stored query aligns perfectly with itself.
        hits = pd.DataFrame(
            {"header": ["q", "b", "a", "c"], "norm_score": [9.0, 3.0, 2.0, 1.0]}
        )
        filtered = Embedding_SSEARCH.filter_ranked_hits(hits, "q", manual_query, top_k)
        return filtered["header"].tolist()

    def test_stored_query_drops_its_self_hit_and_uses_one_top_k_slot(self):
        self.assertEqual(self.ranked(manual_query=False, top_k=3), ["b", "a"])
        self.assertEqual(self.ranked(manual_query=False, top_k=1), [])
        self.assertEqual(self.ranked(manual_query=False, top_k=0), [])
        self.assertEqual(
            self.ranked(manual_query=False, top_k=None), ["b", "a", "c"]
        )

    def test_manual_query_keeps_a_same_header_record_and_all_top_k_slots(self):
        self.assertEqual(self.ranked(manual_query=True, top_k=3), ["q", "b", "a"])
        self.assertEqual(
            self.ranked(manual_query=True, top_k=None), ["q", "b", "a", "c"]
        )


class EmbeddingSsearchMetadataReportTests(unittest.TestCase):
    def test_xlsx_matches_metadata_template_and_imports_into_viewer(self):
        results = pd.DataFrame(
            [
                {
                    "index": -1,
                    "header": "(Query) query_node",
                    "raw_score": 0.0,
                    "norm_score": 99.9,
                    "length": 90,
                    "seq_len": 90,
                    "aln_len": 90,
                    "identity": 100.0,
                },
                {
                    "index": 4,
                    "header": "node_alpha",
                    "raw_score": 42.5,
                    "norm_score": 0.625,
                    "length": 100,
                    "seq_len": 100,
                    "aln_len": 85,
                    "identity": 100.0 * 36 / 85,
                },
                {
                    "index": 7,
                    "header": "node_beta",
                    "raw_score": 37.25,
                    "norm_score": 0.5,
                    "length": 120,
                    "seq_len": 120,
                    "aln_len": 90,
                    "identity": 25.0,
                },
            ]
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            with (
                mock.patch.object(Embedding_SSEARCH, "REPORT_DIR", temp_dir),
                mock.patch.object(Embedding_SSEARCH, "INPUT_EMBED", "db.h5"),
                mock.patch.object(Embedding_SSEARCH, "TOP_K", 10),
                mock.patch.object(Embedding_SSEARCH, "NORM_THRESHOLD", None),
                mock.patch.object(Embedding_SSEARCH, "ALIGNMENT_MODE", "local"),
                mock.patch.object(Embedding_SSEARCH, "GENERATE_FASTA", False),
                redirect_stdout(StringIO()),
            ):
                Embedding_SSEARCH.save_results(
                    results,
                    ("query_node", 90),
                    3,
                    {},
                    "metadata_test",
                    "A" * 90,
                    "longer_sequence",
                    -2.0,
                )

            output_path = os.path.join(temp_dir, "Report_metadata_test.xlsx")
            self.assertTrue(os.path.isfile(output_path))
            with open(
                os.path.join(temp_dir, "Report_metadata_test.txt"),
                encoding="utf-8",
            ) as handle:
                report_text = handle.read()
            self.assertIn(
                " Identity:    Identical standard residues / alignment "
                "length incl. internal gaps",
                report_text,
            )
            self.assertIn("| ALN-LEN  | IDENT% | HEADER", report_text)
            self.assertIn("| 85       | 42.4   | node_alpha", report_text)
            self.assertIn("| 90       | 25.0   | node_beta", report_text)

            raw = pd.read_excel(
                output_path,
                sheet_name="Search Results",
                header=None,
            )
            self.assertEqual(
                raw.iloc[0].tolist(),
                [
                    "Node ID",
                    "Rank",
                    "Norm_Score",
                    "Raw_Score",
                    "Sequence_Length",
                    "Alignment_Length",
                    "Percent_Identity",
                ],
            )
            self.assertEqual(
                raw.iloc[1].tolist(),
                ["Data Type"] + ["number"] * 6,
            )
            self.assertEqual(raw.iloc[2, 0], "node_alpha")
            self.assertEqual(raw.iloc[3, 0], "node_beta")
            self.assertNotIn("(Query) query_node", raw.iloc[:, 0].tolist())

            workbook = load_workbook(output_path, read_only=False)
            self.assertEqual(workbook.sheetnames, ["Search Results", "Search Parameters"])
            worksheet = workbook["Search Results"]
            self.assertEqual(worksheet["A1"].fill.fgColor.rgb, "002C3E50")
            self.assertEqual(worksheet["A2"].fill.fgColor.rgb, "00D5D8DC")
            self.assertEqual(worksheet.freeze_panes, "A3")
            self.assertEqual(worksheet["C3"].number_format, "0.000")
            self.assertEqual(worksheet["G3"].number_format, "0.0")
            self.assertAlmostEqual(worksheet["G3"].value, 100.0 * 36 / 85)
            parameters = workbook["Search Parameters"]
            self.assertEqual(parameters["A1"].fill.fgColor.rgb, "002C3E50")
            self.assertEqual(parameters.column_dimensions["B"].width, 80)
            parameter_values = {
                row[0]: row[1]
                for row in parameters.iter_rows(min_row=2, values_only=True)
            }
            self.assertEqual(
                parameter_values["Percent Identity"],
                "Identical standard residues / alignment length incl. "
                "internal gaps",
            )
            workbook.close()

            viewer = SimpleNamespace(
                full_headers=["node_alpha", "node_beta"],
                n_nodes=2,
                metadata={},
                visible_mask=np.ones(2, dtype=bool),
                selected_indices=[],
                _save_state=mock.Mock(),
                broadcast_event=mock.Mock(),
                get_serializable_metadata=lambda: {},
            )
            with mock.patch.object(Command_Engine, "print_help") as print_help:
                meta_backend.upload_metadata(viewer, [output_path])

            self.assertEqual(
                list(viewer.metadata),
                [
                    "Rank",
                    "Norm_Score",
                    "Raw_Score",
                    "Sequence_Length",
                    "Alignment_Length",
                    "Percent_Identity",
                ],
            )
            self.assertTrue(
                all(entry["type"] == "number" for entry in viewer.metadata.values())
            )
            np.testing.assert_allclose(viewer.metadata["Rank"]["values"], [1, 2])
            np.testing.assert_allclose(
                viewer.metadata["Norm_Score"]["values"],
                [0.625, 0.5],
            )
            np.testing.assert_allclose(
                viewer.metadata["Percent_Identity"]["values"],
                [100.0 * 36 / 85, 25.0],
            )
            print_help.assert_called_once()

    def test_global_report_names_its_end_gap_denominator(self):
        results = pd.DataFrame(
            [
                {
                    "index": 2,
                    "header": "node_gamma",
                    "raw_score": 12.0,
                    "norm_score": 0.1,
                    "length": 120,
                    "seq_len": 120,
                    "aln_len": 120,
                    "identity": 30.0,
                },
            ]
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            with (
                mock.patch.object(Embedding_SSEARCH, "REPORT_DIR", temp_dir),
                mock.patch.object(Embedding_SSEARCH, "ALIGNMENT_MODE", "global"),
                mock.patch.object(Embedding_SSEARCH, "GENERATE_FASTA", False),
                redirect_stdout(StringIO()),
            ):
                Embedding_SSEARCH.save_results(
                    results,
                    ("query_node", 100),
                    1,
                    {},
                    "global_test",
                    "A" * 100,
                    "alignment_length",
                    0.0,
                )
            with open(
                os.path.join(temp_dir, "Report_global_test.txt"),
                encoding="utf-8",
            ) as handle:
                report_text = handle.read()
        self.assertIn(
            " Identity:    Identical standard residues / full alignment "
            "length incl. end gaps",
            report_text,
        )
        self.assertIn("| 120      | 30.0   | node_gamma", report_text)

    def test_text_report_and_hit_fasta_are_written_as_utf8(self):
        # Neither header fits cp1252, the default text encoding of a Windows
        # PC without the system-wide UTF-8 option. Text files opened without
        # an encoding get that default here, whatever this PC uses.
        def open_with_windows_default(file, mode="r", *args, **kwargs):
            if "b" not in mode and len(args) < 2 and kwargs.get("encoding") is None:
                kwargs["encoding"] = "cp1252"
            return open(file, mode, *args, **kwargs)

        header = "node_α_酶"
        results = pd.DataFrame(
            [
                {
                    "index": 2,
                    "header": header,
                    "raw_score": 12.0,
                    "norm_score": 0.1,
                    "length": 120,
                    "seq_len": 120,
                    "aln_len": 120,
                    "identity": 30.0,
                },
            ]
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            with (
                mock.patch.object(Embedding_SSEARCH, "REPORT_DIR", temp_dir),
                mock.patch.object(Embedding_SSEARCH, "ALIGNMENT_MODE", "global"),
                mock.patch.object(Embedding_SSEARCH, "GENERATE_FASTA", True),
                mock.patch.object(
                    Embedding_SSEARCH, "open", open_with_windows_default, create=True
                ),
                redirect_stdout(StringIO()),
            ):
                Embedding_SSEARCH.save_results(
                    results,
                    ("query_β", 100),
                    1,
                    {header: "MKV"},
                    "utf8_test",
                    "A" * 100,
                    "alignment_length",
                    0.0,
                )
            with open(os.path.join(temp_dir, "Report_utf8_test.txt"), "rb") as handle:
                report_text = handle.read().decode("utf-8")
            with open(os.path.join(temp_dir, "Hits_utf8_test.fasta"), "rb") as handle:
                hits_text = handle.read().decode("utf-8").replace("\r\n", "\n")
        self.assertIn(header, report_text)
        self.assertIn("query_β", report_text)
        self.assertEqual(hits_text, f">query_β\n{'A' * 100}\n>{header}\nMKV\n")


if __name__ == "__main__":
    unittest.main()
