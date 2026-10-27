# Copyright 2026 Xuebin Feng
# SPDX-License-Identifier: Apache-2.0
"""Stopping a background job part-way (utilities.Job_Cancellation).

The check point does nothing outside a job and raises inside one that was told
to stop, in the thread of that job only; the exception is not an Exception, so
a worker's own error handling cannot take it for a failure; and a file object
whose every write is a check point stops the encoders and writers that cannot
be given check points of their own (a PNG, an SVG figure, a workbook) within a
chunk of their output, leaving nothing to fail on when they are collected.
"""
import gc
import io
import os
import re
import sys
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest import mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
SRC = Path(__file__).resolve().parents[1] / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.image as mpimg  # noqa: E402
import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
from matplotlib.figure import Figure  # noqa: E402

from tests.held_file_fixtures import (  # noqa: E402
    held_file_error,
    hold_without_delete_sharing,
    transient_hold,
)
from utilities import Job_Cancellation, Localization  # noqa: E402
from utilities.Job_Cancellation import (  # noqa: E402
    HELD_FILE_RETRY_DELAYS,
    CancellableFile,
    FileInUseError,
    JobCancelled,
    JobHandle,
    bound_to,
    cancel_requested,
    current_job,
    forget_partial_path,
    leave_partial,
    note_partial_path,
    raise_if_cancelled,
    remove_partial,
    replace_file,
    sweep_leftover_partials,
)
from utilities.Language_Packs import catalog_path  # noqa: E402
from utilities.Localization import Message, display_text  # noqa: E402


class StopAfter(io.BytesIO):
    """A file that asks for the job to stop once it has had ``writes`` writes."""

    def __init__(self, token, writes):
        super().__init__()
        self.token, self.remaining, self.count = token, writes, 0

    def write(self, data):
        self.count += 1
        result = super().write(data)
        self.remaining -= 1
        if self.remaining == 0:
            self.token.request()
        return result


class StopAfterOnDisk:
    """The same for a real file, which the owner closes: it has the file's seek and tell too."""

    def __init__(self, handle, token, writes):
        self.handle, self.token, self.remaining, self.count = handle, token, writes, 0

    def write(self, data):
        self.count += 1
        result = self.handle.write(data)
        self.remaining -= 1
        if self.remaining == 0:
            self.token.request()
        return result

    def __getattr__(self, name):
        return getattr(self.handle, name)


class CheckPointTests(unittest.TestCase):
    def test_outside_a_job_every_call_does_nothing(self):
        self.assertFalse(cancel_requested())
        raise_if_cancelled()

    def test_a_job_told_to_stop_raises_at_its_next_check_point_and_not_before(self):
        token = JobHandle()
        with bound_to(token):
            raise_if_cancelled()
            self.assertFalse(cancel_requested())
            token.request()
            self.assertTrue(cancel_requested())
            with self.assertRaises(JobCancelled):
                raise_if_cancelled()
        raise_if_cancelled()  # the token is unbound again

    def test_the_exception_is_not_an_exception_so_error_handlers_cannot_swallow_it(self):
        self.assertTrue(issubclass(JobCancelled, BaseException))
        self.assertFalse(issubclass(JobCancelled, Exception))
        token = JobHandle()
        token.request()
        swallowed = []
        with self.assertRaises(JobCancelled), bound_to(token):
            try:
                raise_if_cancelled()
            except Exception as error:  # a worker's "skip this subset" or "report the failure"
                swallowed.append(error)

        self.assertEqual(swallowed, [])

    def test_the_token_belongs_to_the_thread_that_bound_it(self):
        token = JobHandle()
        token.request()
        seen = []

        def other_thread():
            seen.append(cancel_requested())
            raise_if_cancelled()

        with bound_to(token):
            worker = threading.Thread(target=other_thread)
            worker.start()
            worker.join(10)
            self.assertTrue(cancel_requested())
        self.assertEqual(seen, [False])

    def test_bound_blocks_nest_and_restore_the_one_outside(self):
        outer, inner = JobHandle(), JobHandle()
        outer.request()
        with bound_to(outer):
            with bound_to(inner):
                self.assertFalse(cancel_requested())
            self.assertTrue(cancel_requested())
        self.assertFalse(cancel_requested())

    def test_the_token_is_unbound_when_the_job_ends_by_an_exception(self):
        token = JobHandle()
        token.request()
        with self.assertRaises(JobCancelled), bound_to(token):
            raise_if_cancelled()
        self.assertFalse(cancel_requested())


class CancellableFileTests(unittest.TestCase):
    def test_writes_pass_through_until_the_job_is_told_to_stop(self):
        token = JobHandle()
        target = io.BytesIO()
        wrapper = CancellableFile(target)
        with bound_to(token):
            self.assertEqual(wrapper.write(b"abc"), 3)
            wrapper.flush()
            token.request()
            with self.assertRaises(JobCancelled):
                wrapper.write(b"def")
        self.assertEqual(target.getvalue(), b"abc")

    def test_it_writes_normally_outside_a_job(self):
        target = io.BytesIO()
        CancellableFile(target).write(b"abc")
        self.assertEqual(target.getvalue(), b"abc")

    def test_after_the_stop_it_does_nothing_more_so_a_half_done_writer_ends_quietly(self):
        token = JobHandle()
        token.request()
        target = io.BytesIO(b"keep")
        target.seek(4)
        wrapper = CancellableFile(target)
        with bound_to(token), self.assertRaises(JobCancelled):
            wrapper.write(b"x")
        target.close()  # the owner closes the file; the writer may still be collected
        self.assertEqual(wrapper.write(b"more"), 4)
        # A position of its own, which moves as a file's would.
        self.assertEqual(wrapper.tell(), 8)
        self.assertEqual(wrapper.seek(2), 2)
        self.assertEqual(wrapper.write(b"xyz"), 3)
        self.assertEqual(wrapper.tell(), 5)
        wrapper.flush()
        wrapper.close()

    def test_it_has_no_file_descriptor_so_an_encoder_writes_through_it(self):
        with self.assertRaises(io.UnsupportedOperation):
            CancellableFile(io.BytesIO()).fileno()
        with tempfile.TemporaryFile() as real:
            with self.assertRaises(io.UnsupportedOperation):
                CancellableFile(real).fileno()

    def test_closing_it_leaves_the_file_to_its_owner(self):
        target = io.BytesIO()
        wrapper = CancellableFile(target)
        wrapper.close()
        self.assertFalse(target.closed)
        self.assertFalse(wrapper.closed)


class JobHandleTests(unittest.TestCase):
    """What a job's worker and the scheduler share: the stop request and its stop functions, partial files, progress."""

    def test_the_stop_functions_a_worker_registered_are_called_once_when_the_job_is_told_to_stop(self):
        handle, calls = JobHandle(), []
        handle.on_cancel(lambda: calls.append("process"))
        handle.on_cancel(lambda: calls.append("second"))
        self.assertEqual(calls, [])
        handle.request()
        handle.request()
        self.assertEqual(calls, ["process", "second"])
        self.assertTrue(handle.requested)

    def test_a_stop_function_registered_after_the_request_is_called_at_once(self):
        handle, calls = JobHandle(), []
        handle.request()
        handle.on_cancel(lambda: calls.append("late"))
        self.assertEqual(calls, ["late"])

    def test_a_stop_function_that_fails_does_not_keep_the_others_from_running(self):
        handle, calls = JobHandle(), []

        def broken():
            raise OSError("process already gone")

        handle.on_cancel(broken)
        handle.on_cancel(lambda: calls.append("second"))
        with mock.patch("builtins.print") as said:
            handle.request()
        self.assertIn("process already gone", str(said.call_args))
        self.assertEqual(calls, ["second"])
        self.assertTrue(handle.requested)

    def test_partial_files_are_kept_in_the_order_named_until_forgotten(self):
        handle = JobHandle()
        handle.note_partial(Path("a.partial"))
        handle.note_partial("b.partial")
        self.assertEqual(handle.partial_paths(), ["a.partial", "b.partial"])
        handle.forget_partial("a.partial")
        handle.forget_partial("never named")
        self.assertEqual(handle.partial_paths(), ["b.partial"])

    def test_progress_and_the_log_path_tell_the_scheduler_only_when_they_change(self):
        told = []
        handle = JobHandle(on_change=lambda: told.append(1))
        handle.progress("Folding 1 of 3")
        handle.progress("Folding 1 of 3")
        handle.set_log_path(Path("run.log"))
        handle.set_log_path("run.log")
        handle.progress(None)
        self.assertEqual(len(told), 3)
        self.assertEqual((handle.progress_text, handle.log_path), (None, "run.log"))

    def test_the_module_functions_reach_the_handle_bound_to_the_thread_and_do_nothing_outside(self):
        note_partial_path("outside.partial")
        forget_partial_path("outside.partial")
        self.assertIsNone(current_job())
        handle = JobHandle()
        with bound_to(handle):
            self.assertIs(current_job(), handle)
            note_partial_path("inside.partial")
            self.assertEqual(handle.partial_paths(), ["inside.partial"])
            forget_partial_path("inside.partial")
        self.assertEqual(handle.partial_paths(), [])
        self.assertIsNone(current_job())


class OpenAtFirstWriteTests(unittest.TestCase):
    """CancellableFile.at opens the file only when something is written, and never once the job was told to stop."""

    def test_nothing_exists_until_the_first_write_and_the_file_is_closed_on_exit(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "image.png")
            with CancellableFile.at(path, "wb") as target:
                target.write(b"")
                target.flush()
                self.assertEqual((target.tell(), os.path.exists(path)), (0, False), "an empty write opens nothing")
                target.write(b"abc")
                self.assertTrue(os.path.exists(path))
                target.write(b"def")
            self.assertTrue(target.closed)
            self.assertEqual(Path(path).read_bytes(), b"abcdef")

    def test_a_zip_writers_seek_to_the_start_opens_nothing_either(self):
        # ZipFile seeks to the start of the file it is given before it writes anything.
        import zipfile

        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "book.xlsx")
            with CancellableFile.at(path, "wb") as target:
                archive = zipfile.ZipFile(target, "w")
                self.assertFalse(os.path.exists(path))
                archive.writestr("a.txt", "x")
                archive.close()
            with zipfile.ZipFile(path) as written:
                self.assertEqual(written.read("a.txt"), b"x")

    def test_a_job_told_to_stop_never_opens_the_file(self):
        handle = JobHandle()
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "image.png")
            with bound_to(handle):
                with self.assertRaises(JobCancelled), CancellableFile.at(path, "wb") as target:
                    handle.request()
                    target.write(b"abc")
            self.assertEqual(os.listdir(directory), [])

    def test_a_text_file_takes_text_and_a_binary_one_bytes_as_files_do(self):
        # matplotlib tells a text file from a binary one by whether a bytes write is refused.
        with tempfile.TemporaryDirectory() as directory:
            with CancellableFile.at(os.path.join(directory, "t.svg"), "w", "utf-8") as text:
                with self.assertRaises(TypeError):
                    text.write(b"")
                text.write("é")
            with CancellableFile.at(os.path.join(directory, "b.png"), "wb") as binary:
                with self.assertRaises(TypeError):
                    binary.write("text")
            self.assertEqual(Path(directory, "t.svg").read_bytes(), "é".encode("utf-8"))


class EncodersThroughItTests(unittest.TestCase):
    """The writers print, logo and label hand it to: the same bytes, and a stop within a chunk."""

    def test_a_png_through_it_is_the_png_imsave_writes_to_a_path(self):
        image = (np.random.default_rng(7).random((60, 80, 4)) * 255).astype(np.uint8)
        with tempfile.TemporaryDirectory() as directory:
            direct = os.path.join(directory, "direct.png")
            mpimg.imsave(direct, image)
            through = os.path.join(directory, "through.png")
            with open(through, "wb") as handle:
                mpimg.imsave(CancellableFile(handle), image, format="png")
            self.assertEqual(Path(direct).read_bytes(), Path(through).read_bytes())

    def test_a_png_encode_stops_within_a_few_chunks_of_the_output(self):
        # Noise does not compress, so the PNG goes out in many 64 kB writes.
        image = (np.random.default_rng(7).random((900, 1200, 4)) * 255).astype(np.uint8)
        token = JobHandle()
        sink = StopAfter(token, 4)
        with bound_to(token), self.assertRaises(JobCancelled):
            mpimg.imsave(CancellableFile(sink), image, format="png")
        complete = io.BytesIO()
        mpimg.imsave(complete, image, format="png")
        self.assertGreater(len(complete.getvalue()), 20 * 65536, "the picture takes many chunks")
        self.assertLessEqual(sink.count, 6, "the encoder stopped within a chunk or two of the request")
        self.assertLess(len(sink.getvalue()), len(complete.getvalue()) // 10)

    def figure(self):
        figure = Figure(figsize=(4, 3))
        axes = figure.subplots()
        axes.plot(*np.random.default_rng(3).random((2, 4000)))
        return figure

    def test_an_svg_figure_through_it_is_the_svg_savefig_writes_to_a_path(self):
        meta = {"Date": None}
        figure = self.figure()
        with tempfile.TemporaryDirectory() as directory, matplotlib.rc_context({"svg.hashsalt": "fixed"}):
            direct = os.path.join(directory, "direct.svg")
            figure.savefig(direct, format="svg", metadata=meta)
            through = os.path.join(directory, "through.svg")
            # Opened as matplotlib opens a path for an SVG: text, so line ends are the platform's.
            with open(through, "w", encoding="utf-8") as handle:
                figure.savefig(CancellableFile(handle), format="svg", metadata=meta)
            self.assertEqual(Path(direct).read_bytes(), Path(through).read_bytes())

    def test_a_png_figure_through_it_is_the_png_savefig_writes_to_a_path(self):
        figure = self.figure()
        with tempfile.TemporaryDirectory() as directory:
            direct = os.path.join(directory, "direct.png")
            figure.savefig(direct, format="png", transparent=True, dpi=120, bbox_inches="tight")
            through = os.path.join(directory, "through.png")
            with open(through, "wb") as handle:
                figure.savefig(CancellableFile(handle), format="png", transparent=True, dpi=120, bbox_inches="tight")
            self.assertEqual(Path(direct).read_bytes(), Path(through).read_bytes())

    def test_an_svg_figure_stops_while_it_is_drawn(self):
        token = JobHandle()
        sink = StopAfter(token, 20)
        with bound_to(token), self.assertRaises(JobCancelled):
            self.figure().savefig(CancellableFile(sink), format="svg")
        whole = io.BytesIO()
        self.figure().savefig(whole, format="svg")
        self.assertLess(len(sink.getvalue()), len(whole.getvalue()) // 2)

    def workbook(self):
        book = openpyxl.Workbook()
        sheet = book.active
        rng = np.random.default_rng(5)
        for row in rng.random((400, 30)):
            sheet.append([float(value) for value in row])
        return book

    def test_a_workbook_through_it_loads_back_as_saved(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "book.xlsx")
            with open(path, "wb") as handle:
                self.workbook().save(CancellableFile(handle))
            loaded = openpyxl.load_workbook(path)
            self.assertEqual(loaded.active.max_row, 400)
            self.assertEqual(loaded.active.max_column, 30)

    def test_a_workbook_stops_while_it_is_saved_and_its_writer_collects_quietly(self):
        # The workbook's ZipFile closes itself when it is collected, after the owner closed the file:
        # that must not raise (an exception there is printed as "Exception ignored" and lost).
        unraisable = []
        original = sys.unraisablehook
        sys.unraisablehook = lambda info: unraisable.append(info)
        self.addCleanup(setattr, sys, "unraisablehook", original)
        token = JobHandle()
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "book.xlsx")
            with bound_to(token), self.assertRaises(JobCancelled):
                with open(path, "wb") as handle:
                    sink = StopAfterOnDisk(handle, token, 3)
                    self.workbook().save(CancellableFile(sink))
            self.assertTrue(handle.closed)
            gc.collect()
            self.assertEqual(unraisable, [])
            self.assertGreaterEqual(sink.count, 3)


class HeldFileTestCase(unittest.TestCase):
    """A folder of its own, and the list of partial files that could not be deleted private to the test."""

    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.directory = Path(directory.name)
        for patch in (
            mock.patch.object(Job_Cancellation, "_leftover_partials", []),
            mock.patch.object(Job_Cancellation, "_leftover_sweep_registered", False),
        ):
            patch.start()
            self.addCleanup(patch.stop)
        registered = mock.patch.object(Job_Cancellation.atexit, "register")
        self.registered = registered.start()
        self.addCleanup(registered.stop)

    def make(self, name, content=b"new"):
        path = self.directory / name
        path.write_bytes(content)
        return str(path)

    @staticmethod
    def hold_always(operation, winerror=32):
        """os.replace or os.remove (the name) fails as Windows does for a file another program holds; the attempts."""
        attempts = []

        def refuse(*paths):
            attempts.append(paths)
            raise held_file_error(winerror, paths[0])

        return mock.patch.object(Job_Cancellation.os, operation, side_effect=refuse), attempts


class ReplaceFileTests(HeldFileTestCase):
    """The rename of a finished partial file over its target waits out another program that holds either."""

    def test_the_waits_add_up_to_about_two_and_a_half_seconds(self):
        self.assertTrue(2.0 <= sum(HELD_FILE_RETRY_DELAYS) <= 3.0, sum(HELD_FILE_RETRY_DELAYS))

    def test_a_hold_that_ends_is_waited_out_whichever_error_windows_gives(self):
        for winerror in (32, 33, 5):
            with self.subTest(winerror=winerror):
                source, target = self.make("a.partial"), self.make("a.png", b"old")
                with transient_hold(3, winerror) as (attempts, pauses):
                    replace_file(source, target)
                self.assertEqual(Path(target).read_bytes(), b"new")
                self.assertFalse(os.path.exists(source))
                self.assertEqual(len(attempts), 4)
                self.assertEqual(pauses, list(HELD_FILE_RETRY_DELAYS[:3]))

    def test_without_a_hold_it_is_one_rename_and_no_wait(self):
        source, target = self.make("b.partial"), self.make("b.png", b"old")
        with transient_hold(0) as (attempts, pauses):
            replace_file(source, target)
        self.assertEqual((len(attempts), pauses), (1, []))
        self.assertEqual(Path(target).read_bytes(), b"new")

    def test_a_hold_that_never_ends_fails_after_the_waits_and_says_another_program_held_the_file(self):
        source, target = self.make("c.partial"), self.make("c.png", b"old")
        with transient_hold(10 ** 6) as (attempts, pauses):
            with self.assertRaises(FileInUseError) as raised:
                replace_file(source, target)
        self.assertEqual(len(attempts), len(HELD_FILE_RETRY_DELAYS) + 1)
        self.assertEqual(pauses, list(HELD_FILE_RETRY_DELAYS))
        error = raised.exception
        self.assertIsInstance(error, PermissionError)
        self.assertIn("another program", str(error))
        self.assertIn("virus scan", str(error))
        self.assertIn(target, str(error))
        self.assertIn("being used by another process", str(error), "the system's own words follow")
        self.assertEqual(error.__cause__.winerror, 32)
        # The earlier file and the finished one are both as they were.
        self.assertEqual(Path(target).read_bytes(), b"old")
        self.assertTrue(os.path.exists(source))

    def test_the_failure_shows_translated_and_keeps_the_systems_error_in_english_text(self):
        source, target = self.make("d.partial"), self.make("d.png", b"old")
        with transient_hold(10 ** 6):
            with self.assertRaises(FileInUseError) as raised:
                replace_file(source, target, delays=())
        error = raised.exception
        self.assertIsInstance(error.args[0], Message)
        previous = Localization.set_translator(lambda template, n: "[zh] " + template)
        try:
            shown = display_text(error)
        finally:
            Localization.set_translator(previous)
        self.assertTrue(shown.startswith("[zh] Could not save"), shown)
        self.assertIn(target, shown)
        self.assertIn("being used by another process", shown)

    def test_the_failure_text_is_in_the_catalogs_and_has_a_chinese_draft_with_its_placeholders(self):
        source, target = self.make("d2.partial"), self.make("d2.png", b"old")
        with transient_hold(10 ** 6):
            with self.assertRaises(FileInUseError) as raised:
                replace_file(source, target, delays=())
        template = raised.exception.args[0].template
        key = (Localization.MESSAGE_CONTEXT, template, "")
        english = {message.key: message for message in Localization.read_catalog(Localization.TEMPLATE_CATALOG)}
        chinese = {message.key: message for message in Localization.read_catalog(catalog_path("zh_CN"))}
        self.assertIn(key, english)
        self.assertIn(key, chinese)
        draft = chinese[key].translations[0]
        self.assertTrue(draft, "a Chinese draft")
        self.assertEqual(chinese[key].status, "unfinished", "a draft awaits review, as the others do")
        self.assertEqual(set(re.findall(r"\{\w+\}", draft)), {"{path}", "{error}"})
        previous = Localization.set_translator(lambda text, n: draft if text == template else text)
        try:
            shown = display_text(raised.exception)
        finally:
            Localization.set_translator(previous)
        self.assertTrue(shown.startswith("无法保存 "), shown)
        self.assertIn(target, shown)
        self.assertIn("being used by another process", shown)

    def test_any_other_error_is_raised_at_once_and_unchanged(self):
        errors = (
            PermissionError(13, "Access is denied"),                      # a permission, not a hold: no WinError
            held_file_error(winerror=1224),                               # a mapped file: not one of the three
            FileNotFoundError(2, "no such file"),
            OSError(28, "No space left on device"),
        )
        for error in errors:
            with self.subTest(error=repr(error)):
                with mock.patch.object(Job_Cancellation, "_on_windows", return_value=True), \
                        mock.patch.object(Job_Cancellation, "_pause") as pause, \
                        mock.patch.object(Job_Cancellation.os, "replace", side_effect=error) as replace, \
                        self.assertRaises(OSError) as raised:
                    replace_file("a", "b")
                self.assertIs(raised.exception, error)
                self.assertEqual(replace.call_count, 1)
                pause.assert_not_called()

    def test_off_windows_a_refusal_is_the_files_own_and_is_not_retried(self):
        error = held_file_error(32)
        with mock.patch.object(Job_Cancellation, "_on_windows", return_value=False), \
                mock.patch.object(Job_Cancellation, "_pause") as pause, \
                mock.patch.object(Job_Cancellation.os, "replace", side_effect=error) as replace, \
                self.assertRaises(PermissionError) as raised:
            replace_file("a", "b")
        self.assertIs(raised.exception, error)
        self.assertNotIsInstance(error, FileInUseError)
        self.assertEqual(replace.call_count, 1)
        pause.assert_not_called()

    def test_the_caller_may_name_its_own_waits(self):
        source, target = self.make("e.partial"), self.make("e.png", b"old")
        with transient_hold(2) as (attempts, pauses):
            replace_file(source, target, delays=(0.2,) * 9)
        self.assertEqual(pauses, [0.2, 0.2])
        with transient_hold(10 ** 6) as (attempts, pauses):
            with self.assertRaises(FileInUseError):
                replace_file(source, target, delays=(0.2,) * 9)
        self.assertEqual((len(attempts), len(pauses)), (10, 9))

    def test_outside_a_job_the_wait_is_a_sleep(self):
        source, target = self.make("f.partial"), self.make("f.png", b"old")
        real_replace, failures = os.replace, []

        def replace(a, b):
            failures.append(1)
            if len(failures) <= 2:
                raise held_file_error(32, a)
            real_replace(a, b)

        self.assertIsNone(current_job())
        with mock.patch.object(Job_Cancellation, "_on_windows", return_value=True), \
                mock.patch.object(Job_Cancellation.time, "sleep") as sleep, \
                mock.patch.object(Job_Cancellation.os, "replace", side_effect=replace):
            replace_file(source, target)
        self.assertEqual([call.args for call in sleep.call_args_list], [(0.05,), (0.1,)])
        self.assertEqual(Path(target).read_bytes(), b"new")

    def test_a_stop_during_a_wait_ends_it_at_once(self):
        source, target = self.make("g.partial"), self.make("g.png", b"old")
        handle = JobHandle()
        refusing, attempts = self.hold_always("replace")
        timer = threading.Timer(0.05, handle.request)
        started = time.monotonic()
        with mock.patch.object(Job_Cancellation, "_on_windows", return_value=True), refusing, bound_to(handle):
            timer.start()
            try:
                with self.assertRaises(JobCancelled):
                    replace_file(source, target, delays=(30.0, 30.0))
            finally:
                timer.cancel()
        self.assertLess(time.monotonic() - started, 10.0, "it did not wait out the 30 seconds")
        self.assertEqual(len(attempts), 1)

    def test_a_job_told_to_stop_already_does_not_wait_at_all(self):
        source, target = self.make("h.partial"), self.make("h.png", b"old")
        handle = JobHandle()
        handle.request()
        refusing, attempts = self.hold_always("replace")
        with mock.patch.object(Job_Cancellation, "_on_windows", return_value=True), refusing, bound_to(handle):
            with self.assertRaises(JobCancelled):
                replace_file(source, target, delays=(30.0,))
        self.assertEqual(len(attempts), 1)


class RemovePartialTests(HeldFileTestCase):
    """The deletion of a partial file after a failure waits out another program too, and never raises."""

    def test_an_existing_file_is_deleted_and_a_missing_one_counts_as_deleted(self):
        partial = self.make("a.partial")
        self.assertTrue(remove_partial(partial))
        self.assertFalse(os.path.exists(partial))
        self.assertTrue(remove_partial(partial))
        self.assertEqual(Job_Cancellation._leftover_partials, [])

    def test_a_hold_that_ends_is_waited_out(self):
        partial = self.make("b.partial")
        real_remove, attempts, pauses = os.remove, [], []

        def remove(path):
            attempts.append(path)
            if len(attempts) <= 3:
                raise held_file_error(32, path)
            real_remove(path)

        with mock.patch.object(Job_Cancellation, "_on_windows", return_value=True), \
                mock.patch.object(Job_Cancellation, "_pause", side_effect=lambda seconds: pauses.append(seconds)), \
                mock.patch.object(Job_Cancellation.os, "remove", side_effect=remove):
            self.assertTrue(remove_partial(partial))
        self.assertFalse(os.path.exists(partial))
        self.assertEqual(pauses, list(HELD_FILE_RETRY_DELAYS[:3]))
        self.assertEqual(Job_Cancellation._leftover_partials, [])

    def test_a_hold_that_never_ends_leaves_the_file_to_be_deleted_later_and_raises_nothing(self):
        partial = self.make("c.partial")
        refusing, attempts = self.hold_always("remove")
        with mock.patch.object(Job_Cancellation, "_on_windows", return_value=True), \
                mock.patch.object(Job_Cancellation, "_pause", return_value=False) as pause, refusing:
            self.assertFalse(remove_partial(partial))
        self.assertEqual(len(attempts), len(HELD_FILE_RETRY_DELAYS) + 1)
        self.assertEqual([call.args[0] for call in pause.call_args_list], list(HELD_FILE_RETRY_DELAYS))
        self.assertEqual(Job_Cancellation._leftover_partials, [partial])
        self.assertTrue(os.path.exists(partial))
        # Once the program lets go, the sweep deletes it.
        self.assertEqual(sweep_leftover_partials(), 0)
        self.assertFalse(os.path.exists(partial))

    def test_any_other_error_is_not_retried_but_the_file_is_still_left_to_be_deleted_later(self):
        partial = self.make("d.partial")
        with mock.patch.object(Job_Cancellation, "_on_windows", return_value=True), \
                mock.patch.object(Job_Cancellation, "_pause") as pause, \
                mock.patch.object(Job_Cancellation.os, "remove", side_effect=PermissionError(13, "denied")) as remove:
            self.assertFalse(remove_partial(partial))
        self.assertEqual(remove.call_count, 1)
        pause.assert_not_called()
        self.assertEqual(Job_Cancellation._leftover_partials, [partial])

    def test_off_windows_it_is_one_try(self):
        partial = self.make("e.partial")
        with mock.patch.object(Job_Cancellation, "_on_windows", return_value=False), \
                mock.patch.object(Job_Cancellation, "_pause") as pause, \
                mock.patch.object(Job_Cancellation.os, "remove", side_effect=held_file_error(32)) as remove:
            self.assertFalse(remove_partial(partial))
        self.assertEqual(remove.call_count, 1)
        pause.assert_not_called()

    def test_a_job_told_to_stop_tries_once_and_does_not_wait(self):
        partial = self.make("f.partial")
        handle = JobHandle()
        handle.request()
        refusing, attempts = self.hold_always("remove")
        started = time.monotonic()
        with mock.patch.object(Job_Cancellation, "_on_windows", return_value=True), refusing, bound_to(handle):
            self.assertFalse(remove_partial(partial))
        self.assertLess(time.monotonic() - started, 1.0)
        self.assertEqual(len(attempts), 1)
        self.assertEqual(Job_Cancellation._leftover_partials, [partial])

    def test_a_stop_during_the_wait_ends_it_and_leaves_the_file_to_be_deleted_later(self):
        partial = self.make("g.partial")
        handle = JobHandle()
        refusing, attempts = self.hold_always("remove")
        timer = threading.Timer(0.05, handle.request)
        started = time.monotonic()
        with mock.patch.object(Job_Cancellation, "_on_windows", return_value=True), refusing, bound_to(handle):
            timer.start()
            try:
                self.assertFalse(remove_partial(partial, delays=(30.0, 30.0)))
            finally:
                timer.cancel()
        self.assertLess(time.monotonic() - started, 10.0)
        self.assertEqual(len(attempts), 1)
        self.assertEqual(Job_Cancellation._leftover_partials, [partial])


class LeftoverPartialTests(HeldFileTestCase):
    """A partial file that stays held is deleted later: at a job's end, when the Viewer closes, or when the process ends."""

    def test_the_sweep_deletes_what_can_be_deleted_now_and_keeps_the_rest_listed(self):
        gone, held, missing = self.make("a.partial"), self.make("b.partial"), str(self.directory / "c.partial")
        for path in (gone, held, missing):
            leave_partial(path)
        real_remove = os.remove

        def remove(path):
            if path == held:
                raise held_file_error(32, path)
            real_remove(path)

        with mock.patch.object(Job_Cancellation.os, "remove", side_effect=remove):
            self.assertEqual(sweep_leftover_partials(), 1)
        self.assertEqual(Job_Cancellation._leftover_partials, [held])
        self.assertFalse(os.path.exists(gone))
        self.assertTrue(os.path.exists(held))
        self.assertEqual(sweep_leftover_partials(), 0)
        self.assertFalse(os.path.exists(held))

    def test_a_file_is_listed_once_and_the_sweep_is_registered_to_run_when_the_process_ends(self):
        path = self.make("d.partial")
        leave_partial(path)
        leave_partial(Path(path))
        leave_partial(self.make("e.partial"))
        self.assertEqual(len(Job_Cancellation._leftover_partials), 2)
        self.registered.assert_called_once_with(sweep_leftover_partials)

    def test_with_nothing_left_the_sweep_does_nothing(self):
        with mock.patch.object(Job_Cancellation.os, "remove") as remove:
            self.assertEqual(sweep_leftover_partials(), 0)
        remove.assert_not_called()


@unittest.skipUnless(sys.platform == "win32", "holds a file open as a Windows virus scanner does")
class RealHeldFileTests(HeldFileTestCase):
    """Windows itself refuses the rename and the deletion of a file another program has open without delete sharing."""

    def test_a_file_held_without_delete_sharing_cannot_be_renamed_replaced_or_deleted(self):
        source, target = self.make("s.partial"), self.make("s.png", b"old")
        with hold_without_delete_sharing(source):
            with self.assertRaises(PermissionError) as held_source:
                os.replace(source, target)
            with self.assertRaises(PermissionError) as held_deletion:
                os.remove(source)
        self.assertEqual((held_source.exception.winerror, held_deletion.exception.winerror), (32, 32))
        with hold_without_delete_sharing(target):
            with self.assertRaises(PermissionError) as held_target:
                os.replace(source, target)
        self.assertEqual(held_target.exception.winerror, 5)

    def test_the_rename_waits_out_a_scan_that_holds_the_finished_file(self):
        source, target = self.make("a.partial"), self.make("a.png", b"old")
        with hold_without_delete_sharing(source, seconds=0.3):
            replace_file(source, target)
        self.assertEqual(Path(target).read_bytes(), b"new")
        self.assertFalse(os.path.exists(source))

    def test_the_rename_waits_out_a_program_that_holds_the_file_it_replaces(self):
        source, target = self.make("b.partial"), self.make("b.png", b"old")
        with hold_without_delete_sharing(target, seconds=0.3):
            replace_file(source, target)
        self.assertEqual(Path(target).read_bytes(), b"new")
        self.assertFalse(os.path.exists(source))

    def test_a_hold_that_outlasts_the_waits_fails_the_rename_and_the_partial_is_deleted_once_it_lets_go(self):
        source, target = self.make("c.partial"), self.make("c.png", b"old")
        with hold_without_delete_sharing(source):
            with self.assertRaises(FileInUseError) as raised:
                replace_file(source, target, delays=(0.01, 0.01))
            self.assertIn("another program", str(raised.exception))
            self.assertEqual(Path(target).read_bytes(), b"old")
            self.assertFalse(remove_partial(source, delays=(0.01,)))
            self.assertEqual(Job_Cancellation._leftover_partials, [source])
            self.assertEqual(sweep_leftover_partials(), 1, "still held")
        self.assertEqual(sweep_leftover_partials(), 0)
        self.assertFalse(os.path.exists(source))

    def test_the_deletion_waits_out_a_scan_too(self):
        partial = self.make("d.partial")
        with hold_without_delete_sharing(partial, seconds=0.3):
            self.assertTrue(remove_partial(partial))
        self.assertFalse(os.path.exists(partial))
        self.assertEqual(Job_Cancellation._leftover_partials, [])


class EveryQueuingCommandCanBeStoppedTests(unittest.TestCase):
    """A command that queues jobs gives each a way to be stopped, or "stop" would only ever say "stopping".

    A job whose worker has no check point finishes whatever it was told, and is reported as completed. This
    holds every command module that queues jobs to name at least one check point, found without naming a command.
    """

    COMMANDS = Path(__file__).resolve().parents[1] / "src" / "commands"
    CHECK_POINTS = ("raise_if_cancelled", "cancel_requested", "CancellableFile", "on_cancel")

    def test_each_command_module_that_queues_a_job_has_a_check_point(self):
        queuing = [path for path in sorted(self.COMMANDS.glob("*.py"))
                   if "scheduler.enqueue(" in path.read_text(encoding="utf-8")]
        self.assertGreaterEqual(len(queuing), 4, "the modules that queue jobs were found")
        for path in queuing:
            with self.subTest(module=path.name):
                source = path.read_text(encoding="utf-8")
                self.assertTrue(any(point in source for point in self.CHECK_POINTS),
                                f"{path.name} queues jobs but names no way to stop one")


if __name__ == "__main__":
    unittest.main()
